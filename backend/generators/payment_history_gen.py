"""In-memory office-reference exports. Never regenerate receipts or upload."""

from datetime import date
from decimal import Decimal
from io import BytesIO
from html import escape

HEADERS = [
    "Payment ID",
    "Receipt date",
    "OR number",
    "Tax years",
    "Penalty (PHP)",
    "Discount (PHP)",
    "Recorded cash (PHP)",
    "Remarks",
]


def _details(section):
    return (
        f"TD {section['td_number']} - Account #{section['property_id']}"
        + (" (selected)" if section["selected"] else " (previous TD)")
        + (" - ARCHIVED" if section["archived"] else "")
    )


def history_pdf(payload):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import (
        SimpleDocTemplate,
        Paragraph,
        Spacer,
        LongTable,
        TableStyle,
    )

    out = BytesIO()
    styles = getSampleStyleSheet()
    styles["Heading2"].keepWithNext = True
    small = ParagraphStyle(
        "HistorySmall", parent=styles["BodyText"], fontSize=8, leading=11, spaceAfter=3
    )
    cell = ParagraphStyle(
        "HistoryCell",
        parent=small,
        fontSize=7.5,
        leading=10,
        spaceAfter=0,
        wordWrap="CJK",
    )
    heading = ParagraphStyle("HistoryHeader", parent=cell, textColor=colors.white)
    money_cell = ParagraphStyle("HistoryMoney", parent=cell, alignment=2)
    metadata_style = ParagraphStyle("HistoryMetadata", parent=small, keepWithNext=True)

    def para(value, style=cell):
        return Paragraph(
            escape(
                str(value if value is not None else "Not recorded"), quote=False
            ).replace("\n", "<br/>"),
            style,
        )

    story = [Paragraph("Property payment history", styles["Title"])]
    scope = (
        "Selected account and verified previous-TD chain (separate sections)"
        if payload["include_previous"]
        else "Selected account only"
    )
    for text in [
        f"Scope: {scope}. All recorded years.",
        f"Generated (UTC): {payload['generated_at_utc']}",
        f"{payload['payment_count']} payment entries. Total recorded cash: PHP {Decimal(payload['recorded_cash_total']):,.2f}.",
        f"Dated coverage: {payload['coverage_start'] or 'None'} to {payload['coverage_end'] or 'None'}. Undated entries: {payload['undated_count']}.",
        *payload["notes"],
    ]:
        story.append(para(text, small))
    for section in payload["sections"]:
        story.extend([Spacer(1, 10), para(_details(section), styles["Heading2"])])
        for text in [
            f"Current stored owner: {section['current_owner']}",
            f"Current location: {section['current_barangay'] or '-'} / {section['current_location'] or '-'}",
            f"PIN: {section['pin'] or '-'}; Lot/Block: {section['lot_number'] or '-'} / {section['block_number'] or '-'}; Classification: {section['classification'] or '-'}",
        ]:
            story.append(para(text, metadata_style))
        if section["payments"]:
            rows = [
                [para(_details(section), heading)] + [""] * 7,
                [para(h, heading) for h in HEADERS],
            ]
            for r in section["payments"]:
                rows.append(
                    [
                        para(v, money_cell if i in (4, 5, 6) else cell)
                        for i, v in enumerate(
                            [
                                r["payment_id"],
                                r["receipt_date"],
                                r["or_number"],
                                r["tax_years"],
                                *(
                                    (
                                        f"{Decimal(r[k]):,.2f}"
                                        if r[k] is not None
                                        else None
                                    )
                                    for k in ("penalty", "discount", "cash")
                                ),
                                r["remarks"],
                            ]
                        )
                    ]
                )
            table = LongTable(
                rows,
                colWidths=[49, 65, 70, 74, 67, 67, 84, 317],
                repeatRows=2,
                hAlign="LEFT",
            )
            table.setStyle(
                TableStyle(
                    [
                        ("SPAN", (0, 0), (-1, 0)),
                        ("BACKGROUND", (0, 0), (-1, 1), colors.HexColor("#24364B")),
                        (
                            "ROWBACKGROUNDS",
                            (0, 2),
                            (-1, -1),
                            [colors.white, colors.HexColor("#F0F4F8")],
                        ),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 5),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                    ]
                )
            )
            story.append(table)
        else:
            story.append(para("No payments recorded for this account.", small))
        story.append(
            para(
                f"Account subtotal: PHP {Decimal(section['recorded_cash_total']):,.2f} ({section['payment_count']} entries)",
                small,
            )
        )

    def footer(canvas, doc):
        canvas.setFont("Helvetica", 8)
        canvas.drawString(
            24,
            15,
            f"MTO office reference - selected account #{payload['selected_property_id']}",
        )
        canvas.drawRightString(818, 15, f"Page {doc.page}")

    SimpleDocTemplate(
        out,
        pagesize=landscape(A4),
        leftMargin=24,
        rightMargin=24,
        topMargin=24,
        bottomMargin=30,
    ).build(story, onFirstPage=footer, onLaterPages=footer)
    return out.getvalue()


def history_excel(payload):
    # Use the application's existing runtime dependency, not an authoring
    # service or a new desktop dependency. Values represent a recorded snapshot.
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    summary = wb.active
    summary.title = "History summary"

    def text(ws, row, col, value):
        c = ws.cell(row, col, str(value if value is not None else "Not recorded"))
        # Explicit string type prevents formulas from ORs, names or remarks.
        c.data_type = "s"
        c.number_format = "@"
        c.alignment = Alignment(horizontal="left", vertical="center")
        return c

    text(summary, 2, 1, "Property payment history")
    text(summary, 4, 1, "Selected account ID")
    text(summary, 4, 2, payload["selected_property_id"])
    text(summary, 5, 1, "Scope")
    text(
        summary,
        5,
        2,
        (
            "Selected + previous TDs, separate sections"
            if payload["include_previous"]
            else "Selected account only"
        ),
    )
    text(summary, 6, 1, "Generated (UTC)")
    text(summary, 6, 2, payload["generated_at_utc"])
    text(summary, 7, 1, "Payment entries")
    summary.cell(7, 2, payload["payment_count"])
    text(summary, 8, 1, "Recorded cash (PHP)")
    summary.cell(8, 2, Decimal(payload["recorded_cash_total"])).number_format = (
        "#,##0.00"
    )
    text(summary, 9, 1, "Dated coverage")
    text(
        summary,
        9,
        2,
        f"{payload['coverage_start'] or 'None'} to {payload['coverage_end'] or 'None'}",
    )
    text(summary, 10, 1, "Undated entries")
    summary.cell(10, 2, payload["undated_count"])
    for i, note in enumerate(payload["notes"], 12):
        text(summary, i, 1, note)
        summary.merge_cells(start_row=i, start_column=1, end_row=i, end_column=8)
        summary.cell(i, 1).alignment = Alignment(wrap_text=True, vertical="center")
        summary.row_dimensions[i].height = 32
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 53
    for section in payload["sections"]:
        ws = wb.create_sheet(f"Account {section['property_id']}")
        text(ws, 2, 1, _details(section))
        ws.merge_cells("A2:H2")
        for i, label, value in [
            (3, "Current stored owner", section["current_owner"]),
            (
                4,
                "Current location",
                f"{section['current_barangay'] or '-'} / {section['current_location'] or '-'}",
            ),
            (
                5,
                "PIN / Lot / Block",
                f"{section['pin'] or '-'} / {section['lot_number'] or '-'} / {section['block_number'] or '-'}",
            ),
        ]:
            text(ws, i, 1, f"{label}: {value}")
            ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=8)
        text(ws, 7, 1, "Recorded cash (PHP)")
        ws.merge_cells("A7:C7")
        ws.cell(7, 4, Decimal(section["recorded_cash_total"])).number_format = (
            "#,##0.00"
        )
        for col, label in enumerate(HEADERS, 1):
            text(ws, 9, col, label)
            ws.cell(9, col).fill = PatternFill("solid", fgColor="24364B")
            ws.cell(9, col).font = Font(
                name="Arial", color="FFFFFF", bold=True, size=10
            )
            ws.cell(9, col).alignment = Alignment(horizontal="center", wrap_text=True)
        ws.row_dimensions[9].height = 32
        for c in ws[9]:
            c.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
        for idx, p in enumerate(section["payments"], 10):
            text(ws, idx, 1, p["payment_id"])
            if p["receipt_date"]:
                ws.cell(idx, 2, date.fromisoformat(p["receipt_date"])).number_format = (
                    "yyyy-mm-dd"
                )
            else:
                text(ws, idx, 2, "Not recorded")
            for col, key in [(3, "or_number"), (4, "tax_years"), (8, "remarks")]:
                text(ws, idx, col, p[key])
            for col, key in [(5, "penalty"), (6, "discount"), (7, "cash")]:
                if p[key] is None:
                    text(ws, idx, col, "Not recorded")
                else:
                    ws.cell(idx, col, Decimal(p[key])).number_format = "#,##0.00"
            for col in range(1, 9):
                ws.cell(idx, col).alignment = Alignment(
                    horizontal="right" if col in (5, 6, 7) else "left",
                    vertical="top",
                    wrap_text=col in (3, 4, 8),
                )
            # Account for long OR identifiers as well as the 500-character
            # remarks field. Never truncate the stored reference for print.
            lines = max(
                1 + len(str(p[key] or "")) // width
                for key, width in (
                    ("or_number", 16),
                    ("tax_years", 18),
                    ("remarks", 55),
                )
            )
            ws.row_dimensions[idx].height = max(30, 15 * lines)
        if not section["payments"]:
            text(ws, 10, 1, "No payments recorded")
        ws.auto_filter.ref = f"A9:H{max(9, 9 + len(section['payments']))}"
        ws.freeze_panes = "C10"
        ws.print_title_rows = "1:9"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        for i, width in enumerate([15, 16, 20, 22, 20, 20, 24, 65], 1):
            ws.column_dimensions[get_column_letter(i)].width = width
    for ws in wb:
        ws.sheet_view.showGridLines = False
        for row in ws:
            for c in row:
                if c.row != 9 or ws is summary:
                    c.font = Font(
                        name="Arial", size=14 if c.row == 2 else 10, bold=c.row == 2
                    )
        ws.sheet_properties.tabColor = "24364B"
    out = BytesIO()
    wb.save(out)
    return out.getvalue()
