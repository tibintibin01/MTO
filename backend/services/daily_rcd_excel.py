"""Daily RCD output in the office's exact printed 100/200 form layout."""

from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from backend.services.rcd_office_form import make_office_form
from utils.rcd_metadata import resolve_rcd_metadata, report_identifier_format
from utils.rcd_remittance import build_remittance_review

MONEY = "#,##0.00;[Red](#,##0.00);0.00"
# Keep the approved physical form unchanged. SEF has six collection/stock rows.
OFFICE_FORM_RANGE_CAPACITY = 6


def _write_text(sheet, address, value):
    cell = sheet[address]
    cell.value = str(value)
    cell.data_type = (
        "s"  # Never interpret an officer/report/OR identifier as a formula.
    )
    cell.number_format = report_identifier_format(str(value))


def _write_amount(sheet, address, value):
    cell = sheet[address]
    cell.value = float(Decimal(value))
    # Retain the office's existing accounting/currency format, including its
    # treatment of zeros and negatives. Previously unused General cells need
    # only an explicit two-decimal numeric format, not a new color scheme.
    if cell.number_format == "General":
        cell.number_format = "#,##0.00"


def generate_daily_rcd_excel(payload, report_number, *, review=None):
    if payload["issues"]:
        raise ValueError(
            "Resolve the daily receipt/booklet review findings before exporting."
        )
    if len(payload["ranges"]) > OFFICE_FORM_RANGE_CAPACITY:
        raise ValueError(
            "The exact office 100/200 forms support six receipt ranges per report. "
            "Office approval of a continuation form is required. No receipts were omitted."
        )
    metadata = resolve_rcd_metadata(
        report_number,
        payload["collector"]["name"],
        payload["report_date"],
        required=True,
    )
    remittance = (
        build_remittance_review(payload, review) if review is not None else None
    )
    stock = remittance["stock"] if remittance else None
    wb = Workbook()
    wb.remove(wb.active)
    for fund, amount_key, total_row, stock_row, summary_total, signed_row in (
        ("100", "basic_cash", 21, 37, 52, 70),
        ("200", "sef_cash", 18, 34, 47, 65),
    ):
        ws = make_office_form(wb, fund)
        ws["E5"] = int(fund)
        # Fixed OR date, not TODAY(): reopening must not redate a past RCD.
        ws["K5"] = metadata.report_date
        _write_text(ws, "E6", metadata.accountable_officer)
        _write_text(ws, "K6", metadata.report_number)
        for row, group in enumerate(payload["ranges"], 12):
            _write_text(ws, f"A{row}", group["form_type"])
            _write_text(ws, f"E{row}", group["from_or"])
            _write_text(ws, f"H{row}", group["to_or"])
            ws[f"K{row}"] = float(Decimal(group[amount_key]))
        ws[f"K{total_row}"] = float(Decimal(payload["totals"][amount_key]))
        if stock is None:
            for row, group in enumerate(payload["ranges"], stock_row):
                _write_text(ws, f"A{row}", group["form_type"])
                ws[f"H{row}"] = group["count"]
                _write_text(ws, f"I{row}", group["from_or"])
                _write_text(ws, f"J{row}", group["to_or"])
        else:
            for row, group in enumerate(stock["rows"], stock_row):
                _write_text(ws, f"A{row}", group["form_type"])
                for key, quantity, first, last in (
                    ("opening", "B", "C", "D"),
                    ("received", "E", "F", "G"),
                    ("issued", "H", "I", "J"),
                    ("ending", "K", "L", "M"),
                ):
                    segment = group[key]
                    ws[f"{quantity}{row}"] = segment["count"] if segment else 0
                    if segment:
                        _write_text(ws, f"{first}{row}", segment["from_or"])
                        _write_text(ws, f"{last}{row}", segment["to_or"])
        # This total is known; cash/check classification and actual remittance
        # are not. Leave those lines, opening/closing balances and receipt
        # acceptance blank, with the draft explanation on the supporting sheet.
        ws[f"E{summary_total}"] = float(Decimal(payload["totals"][amount_key]))
        _write_text(ws, f"A{signed_row}", metadata.accountable_officer)
        ws[f"F{signed_row}"] = metadata.report_date
        if remittance:
            data = remittance["funds"][fund]
            # Positions are the existing office cells; no merge/row/column changes.
            beginning_row = 48 if fund == "100" else 43
            remit_row = 29 if fund == "100" else 26
            check_row = beginning_row + 1
            if data["total_accountability"] is None:
                ws[f"E{summary_total}"] = None
            for address, value in (
                (f"F{beginning_row}", data["opening"]),
                (f"E{summary_total - 2}", data["cash"]),
                (f"E{summary_total - 1}", data["checks"]),
                (f"E{summary_total}", data["total_accountability"]),
            ):
                if value is not None:
                    _write_amount(ws, address, value)
            for row, check in enumerate(remittance["checks"], check_row):
                _write_text(ws, f"H{row}", check["check_number"])
                _write_text(ws, f"J{row}", check["payee"])
                _write_amount(ws, f"M{row}", check[fund])
            _write_amount(ws, f"M{summary_total + 6}", data["checks"])
            if data["remitted"] is not None:
                _write_text(
                    ws, f"A{remit_row}", remittance["details"]["receiving_officer"]
                )
                for address, value in (
                    (f"K{remit_row}", data["cash"]),
                    (f"K{remit_row + 1}", data["checks"]),
                    (f"K{remit_row + 2}", data["remitted"]),
                    (f"E{summary_total + 2}", -Decimal(data["remitted"])),
                    (f"E{summary_total + 6}", data["closing"]),
                ):
                    if value is not None:
                        _write_amount(ws, address, value)
            # Receiver's signature/acknowledgment and bank-deposit evidence are
            # never fabricated by an operator's handover declaration.

    details = wb.create_sheet("Receipt Details")
    details.sheet_view.showGridLines = False
    details.append(
        [
            "OR date",
            date.fromisoformat(payload["report_date"]),
            "Collecting officer",
            metadata.accountable_officer,
        ]
    )
    details.append(
        [
            "Report no.",
            metadata.report_number,
            "Source signature",
            payload["source_signature"],
        ]
    )
    details.append([])
    details.append(
        [
            "OR",
            "Booklet first",
            "Booklet last",
            "100 collected",
            "200 collected",
            "Receipt total",
            "Supporting payment IDs",
        ]
    )
    for item in payload["receipts"]:
        details.append(
            [
                item["or_number"],
                item["booklet_first"],
                item["booklet_last"],
                float(Decimal(item["basic_cash"])),
                float(Decimal(item["sef_cash"])),
                float(Decimal(item["cash"])),
                ", ".join(map(str, item["payment_ids"])),
            ]
        )
    details.append(
        [
            "TOTAL",
            None,
            None,
            float(Decimal(payload["totals"]["basic_cash"])),
            float(Decimal(payload["totals"]["sef_cash"])),
            float(Decimal(payload["totals"]["cash"])),
        ]
    )
    for cells in details:
        for cell in cells:
            cell.font = Font(name="Arial", size=10)
            if isinstance(cell.value, str):
                cell.data_type = "s"
                cell.number_format = report_identifier_format(cell.value)
            if cell.column in (4, 5, 6) and isinstance(cell.value, (int, float)):
                cell.number_format = MONEY
    details["B1"].number_format = "mmm d, yyyy"
    for row in (1, 2):
        for cell in details[row]:
            cell.alignment = Alignment(horizontal="left", vertical="center")
    for cell in details[4]:
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="203B61")
    for cell in details[details.max_row]:
        cell.font = Font(name="Arial", size=10, bold=True)
        cell.border = Border(top=Side(style="thin", color="203B61"))
    for column in range(1, 8):
        details.column_dimensions[get_column_letter(column)].width = (
            24 if column != 7 else 32
        )
    details.freeze_panes = "A5"
    details.auto_filter.ref = f"A4:G{4 + len(payload['receipts'])}"
    note_row = details.max_row + 3
    for offset, note in enumerate(
        (
            "COLLECTIONS DRAFT: The printed 100/200 forms follow the office layout. "
            "Remittance, full receipt-stock accountability and signatures are not yet verified.",
            (
                "Section C: office-log opening + received - paid ORs - listed cancelled/spoiled ORs = ending UNUSED forms. Cancelled copies remain retained for submission, not destroyed or relieved."
                if stock is not None
                else "Section C lists only the recorded ORs for this day. Beginning, received, "
                "ending stock and cancellations must be checked against the issuance log."
            ),
            (
                "Cash/check and handover entries are operator-reviewed report metadata, not "
                "payment edits, receiver acknowledgment or a bank-deposit entry. Blank is not zero."
                if remittance
                else "No actual cash/check mix, opening balance, remittance acceptance or closing "
                "balance is inferred. Blank is not zero. No COA certification is asserted."
            ),
        )
    ):
        row = note_row + offset
        details.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        cell = details.cell(row, 1, note)
        cell.font = Font(name="Arial", size=10)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        details.row_dimensions[row].height = 30
    if remittance:
        row = note_row + 4
        fields = remittance["details"]
        for label, value in (
            ("Cash/check breakdown reviewed", "Yes - operator declaration"),
            (
                "Full collection handover declared",
                "Yes" if fields["full_handover_confirmed"] else "Not confirmed",
            ),
            ("Receiving officer (handover only)", fields["receiving_officer"]),
            ("Actual handover date", fields["handover_date"]),
            ("Planned bank deposit - NOT proof", fields["planned_deposit_date"]),
            ("Remittance review signature", remittance["review_signature"]),
        ):
            _write_text(details, f"A{row}", label)
            if value is not None:
                if "date" in label.lower() or label.startswith("Planned bank"):
                    details.cell(row, 3, date.fromisoformat(value)).number_format = (
                        "mmm d, yyyy"
                    )
                else:
                    _write_text(details, f"C{row}", value)
            row += 1
        row += 1
        for values in (
            [
                "OR",
                "Bank",
                "Check number",
                "Payee",
                "Check portion of OR",
                "Whole check (not additive)",
            ],
        ):
            for col, value in enumerate(values, 1):
                _write_text(details, f"{get_column_letter(col)}{row}", value)
        check_header_row = row
        for entry in fields["checks"]:
            row += 1
            for col, key in enumerate(
                ("or_number", "bank", "check_number", "payee"), 1
            ):
                _write_text(details, f"{get_column_letter(col)}{row}", entry[key])
            details.cell(row, 5, float(Decimal(entry["amount"]))).number_format = MONEY
            details.cell(row, 6, float(Decimal(entry["face_amount"]))).number_format = (
                MONEY
            )
        for cells in details.iter_rows(min_row=note_row + 4, max_col=7):
            for cell in cells:
                cell.font = Font(
                    name="Arial", size=10, bold=cell.row == check_header_row
                )
                if isinstance(cell.value, str):
                    cell.alignment = Alignment(horizontal="left", vertical="center")
    if stock is not None:
        row = details.max_row + 3
        _write_text(details, f"A{row}", "Stock register reference")
        _write_text(details, f"C{row}", stock["details"]["register_reference"])
        row += 1
        _write_text(
            details,
            f"A{row}",
            "Stock quantities (shared 100/200 physical series, not additive)",
        )
        row += 1
        for col, key in enumerate(
            ("opening", "received", "paid", "cancelled", "ending"), 1
        ):
            _write_text(details, f"{get_column_letter(col)}{row}", key.title())
            details.cell(row + 1, col, stock["totals"][key]).number_format = "0"
        row += 3
        for col, label in enumerate(
            (
                "OR",
                "Status",
                "Reason",
                "Log reference",
                "Collection amount",
                "Copies retained",
            ),
            1,
        ):
            _write_text(details, f"{get_column_letter(col)}{row}", label)
        header_row = row
        for item in stock["details"]["cancelled"]:
            row += 1
            for col, key in enumerate(
                ("or_number", "status", "reason", "reference"), 1
            ):
                _write_text(details, f"{get_column_letter(col)}{row}", item[key])
            _write_amount(details, f"E{row}", "0.00")
            _write_text(details, f"F{row}", "Yes - for RCD submission")
        for cells in details.iter_rows(min_row=header_row, max_row=row, max_col=7):
            for cell in cells:
                cell.font = Font(name="Arial", size=10, bold=cell.row == header_row)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            details.row_dimensions[cells[0].row].height = 45
    wb.properties.title = "Daily RPT RCD - collections draft"
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
