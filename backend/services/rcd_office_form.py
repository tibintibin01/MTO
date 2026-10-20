"""Blank printed form geometry, transcribed from the office's 100/200 sheets.

The data-only layout contains static form labels and styles, not the supplied
workbook, personal entries, payment values, VBA or hidden calculation columns.
"""

from copy import copy
from functools import lru_cache
import json
from pathlib import Path
from defusedxml.ElementTree import fromstring

from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection
from openpyxl.worksheet.dimensions import SheetFormatProperties
from openpyxl.worksheet.page import PageMargins, PrintOptions, PrintPageSetup
from openpyxl.worksheet.properties import PageSetupProperties


@lru_cache(maxsize=1)
def office_layout():
    path = (
        Path(__file__).resolve().parents[1]
        / "report_templates"
        / "rcd_npdr_layout.json"
    )
    with path.open(encoding="utf-8") as source:
        layout = json.load(source)
    if layout["format_revision"] != 1 or set(layout["sheets"]) != {"100", "200"}:
        raise ValueError("Unsupported office RCD form layout.")
    return layout


def make_office_form(workbook, fund):
    """Create only the printed form. Preserve its rows, merges and print setup."""
    spec = office_layout()["sheets"][fund]
    ws = workbook.create_sheet(f"REMITTANCE {fund}")
    classes = (Font, PatternFill, Border, Alignment, Protection)
    kinds = ("font", "fill", "border", "alignment", "protection")
    components = {
        kind: [
            cls.from_tree(fromstring(value)) for value in spec["style_components"][kind]
        ]
        for kind, cls in zip(kinds, classes)
    }

    def apply_style(item, style_id):
        style = spec["styles"][style_id]
        for index, kind in enumerate(kinds):
            setattr(item, kind, copy(components[kind][style[index]]))
        item.number_format = style[5]

    ws.sheet_format = SheetFormatProperties.from_tree(fromstring(spec["sheet_format"]))
    ws.sheet_view.showGridLines = spec["show_gridlines"]
    ws.sheet_view.zoomScale = spec["zoom"]
    for letter, properties in spec["columns"].items():
        dim = ws.column_dimensions[letter]
        if "width" in properties:
            dim.width = float(properties["width"])
        for key in ("hidden", "collapsed", "outlineLevel", "min", "max"):
            if key in properties:
                setattr(dim, key, int(properties[key]))
        if "style_ref" in properties:
            apply_style(dim, properties["style_ref"])
    for number, properties in spec["rows"].items():
        dim = ws.row_dimensions[int(number)]
        if "ht" in properties:
            dim.height = float(properties["ht"])
        for key in ("hidden", "collapsed", "outlineLevel", "thickTop", "thickBot"):
            if key in properties:
                setattr(dim, key, int(properties[key]))
        if "style_ref" in properties:
            apply_style(dim, properties["style_ref"])
    for merged in spec["merges"]:
        ws.merge_cells(merged)
    # Reapply edge styles after merging, including original merged-cell borders.
    for row, references in enumerate(spec["cells"], 1):
        for column, style_id in enumerate(references, 1):
            apply_style(ws.cell(row, column), style_id)
    for address, label in spec["labels"].items():
        # Original spacer strings are empty on paper, never a numeric zero.
        if label.strip():
            ws[address] = label
    ws.page_setup = PrintPageSetup.from_tree(fromstring(spec["page_setup"]))
    ws.page_margins = PageMargins.from_tree(fromstring(spec["page_margins"]))
    ws.print_options = PrintOptions.from_tree(fromstring(spec["print_options"]))
    ws.sheet_properties.pageSetUpPr = PageSetupProperties.from_tree(
        fromstring(spec["setup_properties"])
    )
    ws.print_area = spec["print_area"]
    return ws
