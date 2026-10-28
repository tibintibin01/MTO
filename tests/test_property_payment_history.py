"""Synthetic office-reference reports; no production DB, vault or cloud."""

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch
import os
import time

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, event, select
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from backend.database import Base
from backend.models import (
    Payment,
    PaymentBilling,
    Property,
    PropertyBilling,
    ReceiptHistory,
)
from backend.services.property_payment_history_service import (
    HistoryScopeError,
    _money,
    _receipt_date,
    get_property_payment_history,
    previous_account_scope,
)
from backend.generators.payment_history_gen import history_excel, history_pdf
from ui.payment_history import (
    PaymentHistoryDialog,
    history_save_options,
    selected_history_account,
)


@pytest.fixture
def db():
    engine = create_engine(
        "sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool
    )
    Base.metadata.create_all(engine)
    with Session(engine) as session:
        yield session
    engine.dispose()


def prop(
    db,
    pid=1,
    *,
    td="TEST-TD",
    previous=None,
    previous_td=None,
    year="2026",
    archived=False,
):
    p = Property(
        id=pid,
        td_number=td,
        owner_name="SYNTHETIC OWNER",
        barangay="TEST BARANGAY",
        location="Synthetic location",
        pin="TEST-PIN",
        lot_number="TEST-LOT",
        effectivity_date=year,
        tax_year=year,
        previous_property_id=previous,
        prev_td_number=previous_td,
        deleted_at=datetime(2026, 1, 1) if archived else None,
        assessed_value=999999,
    )
    db.add(p)
    db.flush()
    return p


def pay(
    db,
    pid,
    payment_id,
    *,
    cash="100.00",
    paid=datetime(2024, 3, 1),
    years="2024",
    or_number="0000123",
    discount="0.00",
    penalty="0.00",
    remarks=None,
):
    p = Payment(
        id=payment_id,
        property_id=pid,
        amount=Decimal(cash),
        date_paid=paid,
        tax_year=years,
        or_number=or_number,
        penalty=Decimal(penalty),
        discount=Decimal(discount),
        remarks=remarks,
        created_at=datetime(2026, 10, 28),
    )
    db.add(p)
    db.flush()
    return p


def collect(db, pid=1, previous=False):
    return get_property_payment_history(pid, include_previous=previous, db_session=db)


def test_exact_id_all_years_date_order_and_no_pdf_or_allocation_double_count(db):
    prop(db)
    prop(db, 2)  # same TD and owner, different account
    pay(db, 1, 1, cash="576.00", years="2023,2024", discount="144.00")
    pay(
        db,
        1,
        2,
        cash="648.00",
        paid=datetime(2026, 1, 5),
        years="2026",
        discount="72.00",
    )
    pay(db, 2, 3, cash="777.00")
    for rid in (1, 2):
        db.add(
            ReceiptHistory(
                id=rid,
                property_id=1,
                payment_id=1,
                file_path=f"synthetic-{rid}.pdf",
                generated_at=datetime(2026, 1, 1),
            )
        )
    for bid, year in [(1, 2023), (2, 2024)]:
        db.add(
            PropertyBilling(id=bid, property_id=1, tax_year=year, assessed_value=1000)
        )
        db.add(
            PaymentBilling(payment_id=1, billing_id=bid, tax_year=year, amount_paid=288)
        )
    db.commit()
    statements = []

    def observe(conn, cursor, statement, parameters, context, executemany):
        statements.append(statement.strip().split()[0].upper())

    event.listen(db.bind, "before_cursor_execute", observe)
    before = list(db.execute(select(Payment.id, Payment.amount, Payment.discount)))
    result = collect(db)
    assert result["payment_count"] == 2
    assert result["recorded_cash_total"] == "1224.00"
    assert result["coverage_start"] == "2024-03-01"  # not created_at
    assert result["coverage_end"] == "2026-01-05"
    rows = result["sections"][0]["payments"]
    assert [r["payment_id"] for r in rows] == [1, 2]
    assert rows[0]["tax_years"] == "2023,2024"
    assert rows[0]["basic_cash"] is None and rows[0]["sef_cash"] is None
    assert before == list(
        db.execute(select(Payment.id, Payment.amount, Payment.discount))
    )
    assert not {"INSERT", "UPDATE", "DELETE", "REPLACE"}.intersection(statements)


def test_different_payments_same_or_are_not_deduplicated(db):
    prop(db)
    pay(db, 1, 1, years="2023")
    pay(db, 1, 2, years="2024")
    assert collect(db)["payment_count"] == 2


def test_unpaginated_history_zero_payments_and_missing_dates(db):
    prop(db)
    empty = collect(db)
    assert empty["payment_count"] == 0 and empty["recorded_cash_total"] == "0.00"
    for i in range(1, 122):
        pay(db, 1, i, or_number=str(i), paid=None if i == 1 else datetime(2026, 1, 1))
    result = collect(db)
    assert result["payment_count"] == 121 and result["undated_count"] == 1
    assert result["sections"][0]["payments"][-1]["payment_id"] == 1


def test_optional_previous_accounts_separate_archived_sections_and_not_descendants(db):
    prop(db, 1, td="OLD", year="2023", archived=True)
    prop(db, 2, td="NEW", previous=1, previous_td="OLD", year="2027")
    prop(db, 3, td="FUTURE", previous=2, previous_td="NEW", year="2029")
    pay(db, 1, 1)
    pay(db, 2, 2, cash="200")
    pay(db, 3, 3, cash="900")
    assert collect(db, 2)["recorded_cash_total"] == "200.00"
    result = collect(db, 2, True)
    assert [s["property_id"] for s in result["sections"]] == [1, 2]
    assert result["sections"][0]["archived"]
    assert result["recorded_cash_total"] == "300.00"
    assert len(collect(db, 1, True)["sections"]) == 1


def graph(pid, previous=None, previous_td=None, td=None, year="2026"):
    return dict(
        id=pid,
        previous_property_id=previous,
        prev_td_number=previous_td,
        td_number=td or f"TD-{pid}",
        effectivity_date=year,
        tax_year=year,
    )


@pytest.mark.parametrize(
    "rows,selected",
    [
        ([graph(1, 2), graph(2, 1)], 2),
        ([graph(1), graph(2, 1), graph(3, 1)], 2),
        ([graph(1), graph(2, 1, "WRONG")], 2),
        ([graph(1), graph(2, 999)], 2),
        ([graph(1, td="SAME"), graph(2, td="SAME"), graph(3, previous_td="SAME")], 3),
        ([graph(1, year="2027"), graph(2, 1, year="2026")], 2),
        ([graph(1, year=None), graph(2, 1)], 2),
        ([graph(1), graph(2, previous_td="EXTERNAL")], 2),
        ([graph(1), graph(2, 1), graph(3, 1, "WRONG")], 2),
    ],
)
def test_unsafe_lineage_blocks_expansion(rows, selected):
    with pytest.raises(HistoryScopeError):
        previous_account_scope(rows, selected)


def test_unique_legacy_link_and_explicit_id_with_duplicate_td():
    assert previous_account_scope(
        [graph(1, td="OLD"), graph(2, previous_td=" old ")], 2
    ) == [1, 2]
    assert previous_account_scope(
        [
            graph(1, td="SAME"),
            graph(2, td="SAME"),
            graph(3, previous=1, previous_td="SAME"),
        ],
        3,
    ) == [1, 3]


def test_bad_lineage_does_not_block_selected_only(db):
    prop(db, previous=999)
    result = collect(db)
    assert not result["previous_available"]
    assert "Missing" in result["previous_scope_note"]
    with pytest.raises(HistoryScopeError):
        collect(db, previous=True)


def test_snapshot_fingerprint_changes_with_payment_or_metadata_not_capture_time(db):
    p = prop(db)
    payment = pay(db, 1, 1)
    a = collect(db)
    assert a["fingerprint"] == collect(db)["fingerprint"]
    payment.remarks = "Recorded correction reference"
    db.commit()
    assert a["fingerprint"] != collect(db)["fingerprint"]
    b = collect(db)
    p.owner_name = "CHANGED SYNTHETIC OWNER"
    db.commit()
    assert b["fingerprint"] != collect(db)["fingerprint"]


def test_routes_auth_validation_snapshot_and_both_formats(db):
    from backend.routes.payments import router
    from backend.deps import get_current_user, get_db

    prop(db)
    pay(db, 1, 1)
    db.commit()
    app = FastAPI()
    app.include_router(router)
    app.dependency_overrides[get_db] = lambda: db
    client = TestClient(app)
    path = "/payments/property/1/history"
    assert client.get(path).status_code == 401
    assert (
        client.get(
            path + "/export", params={"format": "pdf", "fingerprint": "a" * 64}
        ).status_code
        == 401
    )
    app.dependency_overrides[get_current_user] = lambda: {
        "username": "synthetic-viewer",
        "role": "viewer",
    }
    preview = client.get(path)
    assert preview.status_code == 200 and preview.headers["cache-control"] == "no-store"
    f = preview.json()["fingerprint"]
    for format, signature in [("pdf", b"%PDF"), ("xlsx", b"PK")]:
        res = client.get(path + "/export", params={"format": format, "fingerprint": f})
        assert res.status_code == 200 and res.content.startswith(signature)
        assert res.headers["cache-control"] == "no-store"
        assert res.headers["content-disposition"].endswith(f'.{format}"')
    assert (
        client.get(
            path + "/export", params={"format": "pdf", "fingerprint": "a" * 64}
        ).status_code
        == 409
    )
    assert (
        client.get(
            path + "/export", params={"format": "zip", "fingerprint": f}
        ).status_code
        == 422
    )
    assert client.get(path + "/export", params={"format": "pdf"}).status_code == 422
    assert client.get("/payments/property/0/history").status_code == 422
    assert client.get("/payments/property/999/history").status_code == 404


def test_excel_literal_text_numeric_money_dates_and_basic_sef_limitation(db):
    p = prop(db)
    p.owner_name = '=HYPERLINK("https://invalid.example")'
    pay(db, 1, 1, cash="648", or_number="0000123", discount="72", remarks="=1+1")
    payload = collect(db)
    book = load_workbook(BytesIO(history_excel(payload)))
    ws = book["Account 1"]
    assert ws["A3"].data_type == "s" and "=HYPERLINK(" in ws["A3"].value
    assert ws["H10"].data_type == "s" and ws["H10"].value == "=1+1"
    assert ws["C10"].value == "0000123"
    assert ws["G10"].data_type == "n" and ws["G10"].value == 648
    assert ws["F10"].value == 72 and ws["B10"].value.date() == date(2024, 3, 1)
    assert ws.freeze_panes == "C10" and ws.auto_filter.ref == "A9:H10"
    assert "not inferred" in " ".join(
        c.value
        for row in book["History summary"]
        for c in row
        if isinstance(c.value, str)
    )
    assert history_pdf(payload).startswith(b"%PDF")


def test_dialog_defaults_selected_scope_and_cancel_no_download():
    assert selected_history_account({"property_id": 3}) == 3
    assert selected_history_account({"multiple": True}) is None
    assert selected_history_account({"multiple": True}, {"property_id": 4}) == 4
    assert selected_history_account({"property_id": "TD-1"}) is None
    for format in ("pdf", "xlsx"):
        opts = history_save_options(format, 3)
        assert opts["defaultextension"] == f".{format}"
        assert opts["filetypes"][0][1] == f"*.{format}"
    dialog = SimpleNamespace(
        busy=False,
        payload={"fingerprint": "a" * 64, "include_previous": False},
        property_id=3,
    )
    with (
        patch("ui.payment_history.filedialog.asksaveasfilename", return_value=""),
        patch(
            "ui.payment_history.payment_service.download_property_payment_history"
        ) as download,
    ):
        PaymentHistoryDialog.export(dialog, "xlsx")
        download.assert_not_called()


def test_dialog_wrong_extension_and_busy_no_download():
    dialog = SimpleNamespace(
        busy=False,
        payload={"fingerprint": "a" * 64, "include_previous": False},
        property_id=3,
    )
    with (
        patch(
            "ui.payment_history.filedialog.asksaveasfilename", return_value="wrong.pdf"
        ),
        patch("ui.payment_history.messagebox.showwarning") as warning,
        patch(
            "ui.payment_history.payment_service.download_property_payment_history"
        ) as download,
    ):
        PaymentHistoryDialog.export(dialog, "xlsx")
        warning.assert_called_once()
        download.assert_not_called()
    dialog.busy = True
    with patch("ui.payment_history.filedialog.asksaveasfilename") as ask:
        PaymentHistoryDialog.export(dialog, "pdf")
        ask.assert_not_called()


def test_client_requests_get_with_exact_id_and_scope():
    from api_clients import payment_service

    with patch.object(payment_service, "api_request", return_value={}) as request:
        payment_service.get_property_payment_history(41)
        request.assert_called_once_with(
            "GET",
            "/payments/property/41/history",
            params={"include_previous": False},
            queue_offline=False,
        )
    with patch(
        "api_clients.api_helper.api_download_file", return_value="synthetic.xlsx"
    ) as download:
        payment_service.download_property_payment_history(41, True, "xlsx", "a" * 64)
        assert download.call_args.args == (
            "GET",
            "/payments/property/41/history/export",
        )
        assert download.call_args.kwargs["params"]["include_previous"] is True


@pytest.mark.skipif(os.name != "nt", reason="Real Windows Tk layout test")
def test_native_dialog_fixed_account_minimum_layout_and_safe_close(db):
    import customtkinter as ctk

    prop(db)
    pay(db, 1, 1)
    payload = collect(db)
    root = ctk.CTk()
    root.withdraw()
    try:
        with patch(
            "ui.payment_history.payment_service.get_property_payment_history",
            return_value=payload,
        ):
            dialog = PaymentHistoryDialog(root, 1)
            dialog.geometry("850x620")
            deadline = time.monotonic() + 3
            while dialog.payload is None and time.monotonic() < deadline:
                root.update()
                time.sleep(0.01)
            assert dialog.payload is not None
            assert dialog.property_id == 1 and dialog.previous.get() is False
            assert dialog.pdf_btn.cget("state") == "normal"
            assert len(dialog.tree.get_children()) == 1
            root.update()
            for widget in (
                dialog.pdf_btn,
                dialog.excel_btn,
                dialog.previous_check,
                dialog.reload_btn,
            ):
                assert widget.winfo_ismapped()
                assert (
                    widget.winfo_rootx() + widget.winfo_width()
                    <= dialog.winfo_rootx() + dialog.winfo_width()
                )
            dialog.close()
            root.update()
    finally:
        root.destroy()


def test_date_and_datetime_driver_values_and_unknown_date():
    assert _receipt_date(date(2026, 10, 18)) == "2026-10-18"
    assert _receipt_date(datetime(2026, 10, 18, 23, 59)) == "2026-10-18"
    assert _receipt_date(None) is None
    with pytest.raises(HistoryScopeError):
        _receipt_date("not a date")


def test_missing_adjustments_are_not_zero_and_invalid_cash_is_not_hidden():
    assert _money(None, optional=True) is None
    assert _money(0, optional=True) == "0.00"
    for invalid in (None, "NaN", "Infinity", "-1.00", "0.001"):
        with pytest.raises(HistoryScopeError):
            _money(invalid)


def test_closed_dialog_does_not_poll_or_touch_widgets():
    from unittest.mock import MagicMock

    dialog = SimpleNamespace(closed=True, after=MagicMock(), results=MagicMock())
    PaymentHistoryDialog.poll(dialog)
    dialog.after.assert_not_called()
    dialog.results.get_nowait.assert_not_called()


def test_wrong_account_preview_disables_export_and_allows_reload():
    from queue import Queue
    from unittest.mock import MagicMock

    results = Queue()
    results.put(("load", {"selected_property_id": 999}, None))
    dialog = SimpleNamespace(
        closed=False,
        property_id=1,
        payload=None,
        results=results,
        previous=SimpleNamespace(get=lambda: False),
        status=MagicMock(),
        render=MagicMock(),
        poll=MagicMock(),
        set_busy=MagicMock(),
        after=MagicMock(),
    )
    PaymentHistoryDialog.poll(dialog)
    assert dialog.payload is None
    dialog.render.assert_not_called()
    dialog.set_busy.assert_called_once_with(False)
