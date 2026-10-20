"""Synthetic RPT receipts and booklet logs only; no office data or services."""

from datetime import datetime
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

from fastapi import FastAPI
from fastapi.testclient import TestClient
import openpyxl
import pytest
from sqlalchemy import create_engine, inspect, event
from sqlalchemy.dialects import mysql
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from backend.database import Base
from backend.models import (
    Payment,
    Property,
    TaxPolicy,
    RcdCollectingOfficer,
    RcdReceiptBooklet,
    RcdRegisterControl,
)
from backend.routes import daily_rcd as routes
from backend.services.daily_rcd_service import build_daily_rcd, RcdSnapshotChanged
from backend.services.daily_rcd_excel import generate_daily_rcd_excel
from backend.services.rcd_register_service import (
    ensure_rcd_register_schema,
    register_officer,
    register_booklet,
    read_register,
    RcdRegisterUnavailable,
)
from utils.rcd_booklets import receipt_serial, booklet_range


@pytest.fixture
def db():
    engine = create_engine(
        "sqlite://", poolclass=StaticPool, connect_args={"check_same_thread": False}
    )
    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine)()
    ensure_rcd_register_schema(session)
    session.commit()
    yield session
    session.close()
    engine.dispose()


def setup_books(db):
    first = register_officer(db, "SAMPLE COLLECTOR ONE", "sample_admin")
    second = register_officer(db, "SAMPLE COLLECTOR TWO", "sample_admin")
    a = register_booklet(db, first.id, "000123", "000172", "2026-10-01", "sample_admin")
    b = register_booklet(
        db, second.id, "000173", "000222", "2026-10-01", "sample_admin"
    )
    db.add(
        TaxPolicy(tax_year=2026, basic_rate=Decimal("0.01"), sef_rate=Decimal("0.01"))
    )
    db.commit()
    return first, second, a, b


def payment(
    db,
    number="000123",
    cash="648",
    discount="72",
    paid_on=None,
    encoded_on=None,
    year="2026",
    posted_by="different_encoder",
):
    prop = Property(
        td_number=f"SYNTHETIC-{db.query(Property).count()}",
        owner_name="SAMPLE OWNER",
        assessed_value=36000,
    )
    db.add(prop)
    db.flush()
    result = Payment(
        property_id=prop.id,
        or_number=number,
        tax_year=year,
        amount=Decimal(cash),
        penalty=0,
        discount=Decimal(discount),
        date_paid=paid_on or datetime(2026, 10, 18),
        created_at=encoded_on or datetime(2026, 10, 19),
        posted_by=posted_by,
    )
    db.add(result)
    db.commit()
    return result


@pytest.mark.parametrize(
    "first,last",
    [("000123", "000172"), ("OR-2026-001", "OR-2026-050"), ("009951", "010000")],
)
def test_booklets_use_actual_fifty_receipt_bounds_not_ceiling(first, last):
    _, start, end, _ = booklet_range(first, last)
    assert end - start + 1 == 50


@pytest.mark.parametrize(
    "first,last",
    [
        ("000123", "000173"),
        ("000123", "000122"),
        ("A-001", "B-050"),
        ("001", "50"),
        ("1000000000000000", "1000000000000049"),
    ],
)
def test_invalid_or_conflicting_booklet_formats_are_rejected(first, last):
    with pytest.raises(ValueError):
        booklet_range(first, last)


def test_parser_preserves_real_prefix_and_only_normalizes_padding():
    assert receipt_serial(" OR-2026-000123 ") == ("OR-2026-", 123, 6)
    assert receipt_serial("123") == ("", 123, 3)
    with pytest.raises(ValueError):
        receipt_serial("000123 EXTRA")


def test_register_is_additive_idempotent_and_does_not_change_payments(db):
    def columns():
        return [
            (item["name"], str(item["type"]), item["nullable"], item["default"])
            for item in inspect(db.connection()).get_columns("payments")
        ]

    before = columns()
    ensure_rcd_register_schema(db)
    ensure_rcd_register_schema(db)
    db.commit()
    assert db.query(RcdRegisterControl).count() == 1
    assert db.query(Payment).count() == 0
    assert before == columns()


def test_reads_never_initialize_missing_register(db):
    RcdRegisterControl.__table__.drop(db.connection())
    db.commit()
    before = set(inspect(db.connection()).get_table_names())
    with pytest.raises(RcdRegisterUnavailable):
        read_register(db)
    assert set(inspect(db.connection()).get_table_names()) == before


def test_overlap_rejected_across_collectors_and_padding(db):
    first, second, a, _ = setup_books(db)
    with pytest.raises(ValueError, match="overlaps"):
        register_booklet(db, second.id, "123", "172", "2026-10-02", "sample_admin")
    db.rollback()
    assert db.get(RcdReceiptBooklet, a.id).collector_id == first.id
    assert db.query(RcdReceiptBooklet).count() == 2


def test_mysql_overlap_validation_uses_current_reads_after_register_lock(db):
    first, _, _, _ = setup_books(db)
    statements = []

    def record(connection, clause, multiparams, params, execution_options):
        statements.append(str(clause.compile(dialect=mysql.dialect())))

    event.listen(db.get_bind(), "before_execute", record)
    try:
        register_booklet(db, first.id, "000223", "000272", "2026-10-18", "sample_admin")
        db.commit()
    finally:
        event.remove(db.get_bind(), "before_execute", record)
    assert any(
        "FROM rcd_register_control" in sql and "FOR UPDATE" in sql for sql in statements
    )
    assert any(
        "FROM rcd_receipt_booklets" in sql and "FOR UPDATE" in sql for sql in statements
    )


def test_assignments_and_officer_identity_cannot_be_overwritten(db):
    first, second, a, _ = setup_books(db)
    a.collector_id = second.id
    with pytest.raises(ValueError, match="historical assignments"):
        db.commit()
    db.rollback()
    assert db.get(RcdReceiptBooklet, a.id).collector_id == first.id
    first.name = "REPLACEMENT NAME"
    with pytest.raises(ValueError, match="historical assignments"):
        db.commit()
    db.rollback()
    db.delete(db.get(RcdReceiptBooklet, a.id))
    with pytest.raises(ValueError, match="historical assignments"):
        db.commit()
    db.rollback()


def test_daily_report_follows_or_date_and_booklet_owner_not_encoder(db):
    first, second, _, _ = setup_books(db)
    payment(db, posted_by="SAMPLE COLLECTOR TWO")
    payment(db, "000173", cash="576", discount="144", posted_by="SAMPLE COLLECTOR ONE")
    report = build_daily_rcd(db, first.id, "2026-10-18")
    assert report["status"] == "COLLECTIONS_READY"
    assert report["totals"] == {
        "basic_cash": "324.00",
        "sef_cash": "324.00",
        "cash": "648.00",
    }
    assert report["receipt_count"] == 1
    assert report["other_officer_payment_rows"] == 1
    assert build_daily_rcd(db, first.id, "2026-10-19")["receipt_count"] == 0
    assert build_daily_rcd(db, second.id, "2026-10-18")["totals"]["cash"] == "576.00"
    assert db.query(Payment).count() == 2


def test_next_booklet_same_officer_does_not_merge_book_boundaries(db):
    first, _, _, _ = setup_books(db)
    register_booklet(db, first.id, "000223", "000272", "2026-10-18", "sample_admin")
    db.commit()
    payment(db, "000172")
    payment(db, "000223")
    report = build_daily_rcd(db, first.id, "2026-10-18")
    assert len(report["ranges"]) == 2
    assert report["totals"]["cash"] == "1296.00"


def test_or_day_includes_end_of_day_but_not_next_midnight(db):
    first, _, _, _ = setup_books(db)
    payment(db, "000123", paid_on=datetime(2026, 10, 18, 23, 59, 59, 999999))
    payment(db, "000124", paid_on=datetime(2026, 10, 19, 0, 0))
    report = build_daily_rcd(db, first.id, "2026-10-18")
    assert report["receipt_count"] == 1
    assert report["receipts"][0]["or_number"] == "000123"
    assert not report["issues"]


def test_gaps_are_not_reported_as_issued_receipts(db):
    first, _, _, _ = setup_books(db)
    payment(db, "000123")
    payment(db, "000125")
    result = build_daily_rcd(db, first.id, "2026-10-18")
    assert [(r["from_or"], r["to_or"], r["count"]) for r in result["ranges"]] == [
        ("000123", "000123", 1),
        ("000125", "000125", 1),
    ]


def test_same_receipt_multiple_properties_is_one_or_with_all_cash(db):
    first, _, _, _ = setup_books(db)
    payment(db, cash="100", discount="0")
    payment(db, cash="200", discount="0")
    report = build_daily_rcd(db, first.id, "2026-10-18")
    assert report["receipt_count"] == 1 and report["payment_row_count"] == 2
    assert report["totals"]["cash"] == "300.00"


@pytest.mark.parametrize(
    "number,paid_on,code",
    [
        ("999999", datetime(2026, 10, 18), "BOOKLET_ASSIGNMENT_REQUIRED"),
        ("INVALID", datetime(2026, 10, 18), "INVALID_OR"),
        ("000123", datetime(2026, 9, 30), "OR_BEFORE_BOOKLET_ISSUE"),
    ],
)
def test_unmapped_or_invalid_receipts_block_complete_export(db, number, paid_on, code):
    first, _, _, _ = setup_books(db)
    payment(db, number, paid_on=paid_on)
    report = build_daily_rcd(db, first.id, paid_on.date().isoformat())
    assert report["status"] == "REVIEW_REQUIRED"
    assert code in {item["code"] for item in report["issues"]}
    with pytest.raises(ValueError, match="Resolve"):
        generate_daily_rcd_excel(report, "0007")


def test_same_physical_or_under_different_dates_blocks_both_reports(db):
    first, _, _, _ = setup_books(db)
    payment(db, "000123", paid_on=datetime(2026, 10, 18))
    payment(db, "123", paid_on=datetime(2026, 10, 17))
    for day in ("2026-10-17", "2026-10-18"):
        report = build_daily_rcd(db, first.id, day)
        assert "OR_HAS_DIFFERENT_PAYMENT_DATES" in {
            item["code"] for item in report["issues"]
        }


def test_missing_policy_does_not_invent_an_equal_split(db):
    first, _, _, _ = setup_books(db)
    payment(db, year="2025")
    report = build_daily_rcd(db, first.id, "2026-10-18")
    assert report["issues"][0]["code"] == "RECEIPT_REVIEW_REQUIRED"


def test_stale_preview_is_rejected_after_cash_or_register_changes(db):
    first, _, _, _ = setup_books(db)
    receipt = payment(db)
    original = build_daily_rcd(db, first.id, "2026-10-18")
    receipt.amount = Decimal("650.00")
    db.commit()
    with pytest.raises(RcdSnapshotChanged):
        build_daily_rcd(
            db, first.id, "2026-10-18", expected_signature=original["source_signature"]
        )
    updated = build_daily_rcd(db, first.id, "2026-10-18")
    register_booklet(db, first.id, "000223", "000272", "2026-10-18", "sample_admin")
    db.commit()
    with pytest.raises(RcdSnapshotChanged):
        build_daily_rcd(
            db, first.id, "2026-10-18", expected_signature=updated["source_signature"]
        )


def test_excel_has_two_funds_fixed_date_and_no_assumed_remittance(db):
    first, _, _, _ = setup_books(db)
    payment(db)
    report = build_daily_rcd(db, first.id, "2026-10-18")
    output = generate_daily_rcd_excel(report, "0007/2026")
    wb = openpyxl.load_workbook(output)
    assert wb.sheetnames == ["REMITTANCE 100", "REMITTANCE 200", "Receipt Details"]
    for name, total, remittance, closing in (
        ("REMITTANCE 100", "K21", "K29", "E58"),
        ("REMITTANCE 200", "K18", "K26", "E53"),
    ):
        sheet = wb[name]
        assert sheet["K5"].value == datetime(2026, 10, 18)
        assert sheet["K6"].value == "0007/2026"
        assert sheet["E6"].value == "SAMPLE COLLECTOR ONE"
        assert sheet[total].value == 324
        assert sheet["E12"].value == "000123"
        assert sheet["E12"].number_format == "000000"
        assert sheet[remittance].value is None  # Actual remittance is not inferred.
        assert all(cell.value != "=TODAY()" for row in sheet for cell in row)
        assert sheet[closing].value is None
    assert wb["Receipt Details"]["F5"].value == 648
    assert db.query(Payment).one().amount == Decimal("648.00")


def test_unequal_recorded_rate_split_is_not_forced_identical(db):
    first, _, _, _ = setup_books(db)
    policy = db.query(TaxPolicy).one()
    policy.basic_rate = Decimal("0.02")
    db.commit()
    payment(db, cash="900", discount="0")
    report = build_daily_rcd(db, first.id, "2026-10-18")
    assert report["totals"] == {
        "basic_cash": "600.00",
        "sef_cash": "300.00",
        "cash": "900.00",
    }
    wb = openpyxl.load_workbook(generate_daily_rcd_excel(report, "=LITERAL()"))
    assert wb["REMITTANCE 100"]["K21"].value == 600
    assert wb["REMITTANCE 200"]["K18"].value == 300
    assert wb["REMITTANCE 100"]["K6"].data_type == "s"


@pytest.fixture
def client(db):
    app = FastAPI()
    app.include_router(routes.router, prefix="/reports")
    app.dependency_overrides[routes.admin_only] = lambda: {
        "id": 1,
        "username": "sample_admin",
        "role": "admin",
    }
    app.dependency_overrides[routes.get_db] = lambda: db
    with patch.object(routes, "log_action"), TestClient(app) as client:
        yield client


def test_api_preview_export_and_stale_conflict(db, client):
    first, _, _, _ = setup_books(db)
    item = payment(db)
    params = {"collector_id": first.id, "report_date": "2026-10-18"}
    preview = client.get("/reports/rcd/daily", params=params)
    assert preview.status_code == 200
    exported = client.get(
        "/reports/rcd/daily/excel",
        params={
            **params,
            "report_number": "0007",
            "expected_signature": preview.json()["source_signature"],
        },
    )
    assert exported.status_code == 200
    assert exported.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert (
        openpyxl.load_workbook(BytesIO(exported.content)).sheetnames[0]
        == "REMITTANCE 100"
    )
    item.amount = Decimal("650.00")
    db.commit()
    stale = client.get(
        "/reports/rcd/daily/excel",
        params={
            **params,
            "report_number": "0007",
            "expected_signature": preview.json()["source_signature"],
        },
    )
    assert stale.status_code == 409


def test_failed_audit_rolls_back_register_entry(db):
    with patch.object(
        routes, "log_action", side_effect=RuntimeError("synthetic audit failure")
    ):
        with pytest.raises(Exception) as failure:
            routes.add_officer(
                routes.OfficerInput(name="SAMPLE NEW COLLECTOR"),
                current_user={"id": 1, "username": "sample_admin"},
                db=db,
            )
    assert failure.value.status_code == 500
    assert db.query(RcdCollectingOfficer).count() == 0
    assert db.get(RcdRegisterControl, 1).revision == 0


def test_api_rejects_invalid_booklet_and_provides_no_update_delete(client):
    officer = client.post("/reports/rcd/officers", json={"name": "SAMPLE OFFICER"})
    assert officer.status_code == 200
    invalid = client.post(
        "/reports/rcd/booklets",
        json={
            "collector_id": officer.json()["id"],
            "first_or": "001",
            "last_or": "051",
            "issued_on": "2026-10-01",
        },
    )
    assert invalid.status_code == 422
    assert client.get("/reports/rcd/register").json()["booklets"] == []
    assert client.delete("/reports/rcd/booklets/1").status_code == 404
    assert (
        client.patch("/reports/rcd/booklets/1", json={"collector_id": 2}).status_code
        == 404
    )


def test_new_routes_keep_admin_dependency():
    for route in routes.router.routes:
        assert routes.admin_only in [
            dependency.call for dependency in route.dependant.dependencies
        ]


def test_daily_routes_do_not_allow_unauthenticated_access(db):
    app = FastAPI()
    app.include_router(routes.router, prefix="/reports")
    app.dependency_overrides[routes.get_db] = lambda: db
    with TestClient(app) as client:
        response = client.get(
            "/reports/rcd/daily",
            params={"collector_id": 1, "report_date": "2026-10-18"},
        )
    assert response.status_code == 401


def test_reviewed_export_audits_declarations_but_does_not_edit_money_or_bank(
    db, client
):
    from backend.models import BankDeposit

    first, _, _, _ = setup_books(db)
    item = payment(db)
    before = (item.amount, item.discount, item.date_paid, item.created_at)
    source = build_daily_rcd(db, first.id, "2026-10-18")
    body = {
        "collector_id": first.id,
        "report_date": "2026-10-18",
        "report_number": "0007",
        "expected_signature": source["source_signature"],
        "review": {"cash_check_confirmed": True, "checks": []},
    }
    with patch.object(routes, "log_action") as audit:
        exported = client.post("/reports/rcd/daily/excel", json=body)
    assert exported.status_code == 200
    assert "not receiver/bank acceptance" in audit.call_args.kwargs["action"]
    assert "declared_review" in audit.call_args.kwargs["action"]
    assert (item.amount, item.discount, item.date_paid, item.created_at) == before
    assert db.query(BankDeposit).count() == 0
    bad = {**body, "review": {"cash_check_confirmed": False, "checks": []}}
    assert client.post("/reports/rcd/daily/excel", json=bad).status_code == 422
    item.amount = Decimal("650")
    db.commit()
    assert client.post("/reports/rcd/daily/excel", json=body).status_code == 409


def test_reviewed_export_audit_failure_releases_no_file(db, client):
    first, _, _, _ = setup_books(db)
    payment(db)
    source = build_daily_rcd(db, first.id, "2026-10-18")
    with patch.object(
        routes, "log_action", side_effect=RuntimeError("synthetic audit failure")
    ):
        result = client.post(
            "/reports/rcd/daily/excel",
            json={
                "collector_id": first.id,
                "report_date": "2026-10-18",
                "report_number": "0007",
                "expected_signature": source["source_signature"],
                "review": {"cash_check_confirmed": True, "checks": []},
            },
        )
    assert result.status_code == 500
    assert not result.headers["content-type"].startswith("application/vnd")


def _stock_input(booklet_id, cancelled=None):
    return {
        "confirmed": True,
        "register_reference": "SAMPLE STOCK LOG PAGE 1",
        "opening_ranges": [
            {"booklet_id": booklet_id, "from_or": "000123", "to_or": "000172"}
        ],
        "received_booklet_ids": [],
        "cancelled": cancelled or [],
    }


def _cancelled(number):
    return {
        "or_number": number,
        "status": "CANCELLED",
        "reason": "SAMPLE unused damage",
        "reference": "SAMPLE LOG 2",
        "copies_retained": True,
    }


def test_stock_context_contains_only_registered_officer_booklets_by_or_date(db):
    first, _, a, _ = setup_books(db)
    register_booklet(db, first.id, "000223", "000272", "2026-10-19", "sample_admin")
    db.commit()
    source = build_daily_rcd(db, first.id, "2026-10-18")
    assert [b["id"] for b in source["stock_booklets"]] == [a.id]


@pytest.mark.parametrize("date_value", [datetime(2026, 10, 17), None])
def test_stock_history_blocks_prior_or_undated_usage_in_opening(db, date_value):
    from backend.services.daily_rcd_service import validate_stock_history
    from utils.rcd_stock import build_stock_review

    first, _, book, _ = setup_books(db)
    payment(db)
    prior = payment(db, number="000124", paid_on=datetime(2026, 10, 17))
    prior.date_paid = date_value
    db.commit()
    source = build_daily_rcd(db, first.id, "2026-10-18")
    details = build_stock_review(source, _stock_input(book.id))
    with pytest.raises(ValueError, match="earlier or undated"):
        validate_stock_history(db, source, details)


def test_historical_ending_stock_may_be_used_the_next_day_but_not_cancelled(db):
    from backend.services.daily_rcd_service import validate_stock_history
    from utils.rcd_stock import build_stock_review

    first, _, book, _ = setup_books(db)
    payment(db)
    future = payment(db, number="124", paid_on=datetime(2026, 10, 19))
    source = build_daily_rcd(db, first.id, "2026-10-18")
    details = build_stock_review(source, _stock_input(book.id))
    validate_stock_history(db, source, details)
    cancelled = build_stock_review(
        source, _stock_input(book.id, [_cancelled("000124")])
    )
    with pytest.raises(ValueError, match="saved payment"):
        validate_stock_history(db, source, cancelled)
    assert db.get(Payment, future.id).or_number == "124"


def test_reviewed_stock_export_uses_history_gate_and_audits_zero_collection_or(
    db, client
):
    from backend.models import BankDeposit

    first, _, book, _ = setup_books(db)
    original = payment(db)
    before = (original.amount, original.discount, original.date_paid)
    source = build_daily_rcd(db, first.id, "2026-10-18")
    body = {
        "collector_id": first.id,
        "report_date": "2026-10-18",
        "report_number": "0007",
        "expected_signature": source["source_signature"],
        "review": {
            "cash_check_confirmed": True,
            "stock": _stock_input(book.id, [_cancelled("124")]),
        },
    }
    with patch.object(routes, "log_action") as audit:
        exported = client.post("/reports/rcd/daily/excel", json=body)
    assert exported.status_code == 200
    assert '"or_number": "000124"' in audit.call_args.kwargs["action"]
    assert db.query(Payment).count() == 1
    assert db.query(BankDeposit).count() == 0
    assert (original.amount, original.discount, original.date_paid) == before
    payment(db, number="000124", paid_on=datetime(2026, 10, 19))
    # A new future entry does not alter the selected-day signature, but does
    # conflict with the proposed cancellation. The history gate still blocks it.
    with patch.object(routes, "log_action") as audit:
        blocked = client.post("/reports/rcd/daily/excel", json=body)
    assert blocked.status_code == 422
    assert not audit.called


def test_stock_declaration_is_not_released_when_export_audit_fails(db, client):
    first, _, book, _ = setup_books(db)
    payment(db)
    source = build_daily_rcd(db, first.id, "2026-10-18")
    with patch.object(
        routes, "log_action", side_effect=RuntimeError("synthetic audit failure")
    ):
        result = client.post(
            "/reports/rcd/daily/excel",
            json={
                "collector_id": first.id,
                "report_date": "2026-10-18",
                "report_number": "0007",
                "expected_signature": source["source_signature"],
                "review": {
                    "cash_check_confirmed": True,
                    "stock": _stock_input(book.id),
                },
            },
        )
    assert result.status_code == 500
    assert not result.headers["content-type"].startswith("application/vnd")
