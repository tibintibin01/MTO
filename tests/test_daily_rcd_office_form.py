"""Blank office form geometry and synthetic recorded collections only."""

from copy import deepcopy
from datetime import datetime
from io import BytesIO
from xml.etree.ElementTree import fromstring
import zipfile

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection
from openpyxl.worksheet.page import PageMargins, PrintPageSetup
import pytest

from backend.services.daily_rcd_excel import generate_daily_rcd_excel
from backend.services.rcd_office_form import office_layout


def sample():
    return {
        "issues": [],
        "collector": {"id": 1, "name": "SAMPLE COLLECTOR"},
        "report_date": "2026-10-18",
        "source_signature": "synthetic-signature",
        "ranges": [
            {
                "form_type": "AF 56",
                "from_or": "000123",
                "to_or": "000123",
                "count": 1,
                "basic_cash": "324.00",
                "sef_cash": "324.00",
            }
        ],
        "totals": {"basic_cash": "324.00", "sef_cash": "324.00", "cash": "648.00"},
        "receipts": [
            {
                "or_number": "000123",
                "booklet_first": "000123",
                "booklet_last": "000172",
                "basic_cash": "324.00",
                "sef_cash": "324.00",
                "cash": "648.00",
                "payment_ids": [1],
            }
        ],
    }


@pytest.mark.parametrize("fund,last_row,scale", [("100", 72, 92), ("200", 67, 94)])
def test_original_geometry_labels_print_settings_and_all_cell_styles(
    fund, last_row, scale
):
    sheet = openpyxl.load_workbook(generate_daily_rcd_excel(sample(), "0007"))[
        f"REMITTANCE {fund}"
    ]
    spec = office_layout()["sheets"][fund]
    assert sheet.max_row == last_row
    assert sheet.max_column == 13
    assert sorted(str(r) for r in sheet.merged_cells.ranges) == spec["merges"]
    assert sheet.print_area == f"'REMITTANCE {fund}'!$A$1:$M${last_row}"
    assert sheet.page_setup == PrintPageSetup.from_tree(fromstring(spec["page_setup"]))
    assert sheet.page_margins == PageMargins.from_tree(fromstring(spec["page_margins"]))
    assert sheet.page_setup.paperSize == 5  # Legal, not A4 or Letter.
    assert sheet.page_setup.scale == scale
    for letter, properties in spec["columns"].items():
        assert sheet.column_dimensions[letter].width == float(properties["width"])
        assert sheet.column_dimensions[letter].min == int(properties["min"])
        assert sheet.column_dimensions[letter].max == int(properties["max"])
    assert set(sheet.column_dimensions) == set(spec["columns"])
    for row, properties in spec["rows"].items():
        assert sheet.row_dimensions[int(row)].height == (
            float(properties["ht"]) if "ht" in properties else None
        )
    for address, label in spec["labels"].items():
        assert sheet[address].value == (label if label.strip() else None)
    kinds = ("font", "fill", "border", "alignment", "protection")
    classes = (Font, PatternFill, Border, Alignment, Protection)
    components = {
        kind: [cls.from_tree(fromstring(v)) for v in spec["style_components"][kind]]
        for kind, cls in zip(kinds, classes)
    }
    for letter, properties in spec["columns"].items():
        style = spec["styles"][properties["style_ref"]]
        for index, kind in enumerate(kinds):
            assert (
                getattr(sheet.column_dimensions[letter], kind)
                == components[kind][style[index]]
            )
    for row, references in enumerate(spec["cells"], 1):
        for column, style_id in enumerate(references, 1):
            cell = sheet.cell(row, column)
            style = spec["styles"][style_id]
            for index, kind in enumerate(kinds):
                assert getattr(cell, kind) == components[kind][style[index]], (
                    fund,
                    cell.coordinate,
                    kind,
                )


@pytest.mark.parametrize(
    "fund,total,stock,summary,closing,remittance,ack_amount,signature",
    [
        ("100", "K21", 37, "E52", "E58", "K29", "K68", 70),
        ("200", "K18", 34, "E47", "E53", "K26", "K63", 65),
    ],
)
def test_only_known_values_fill_existing_slots(
    fund, total, stock, summary, closing, remittance, ack_amount, signature
):
    sheet = openpyxl.load_workbook(generate_daily_rcd_excel(sample(), "=LITERAL()"))[
        f"REMITTANCE {fund}"
    ]
    assert sheet["E5"].value == int(fund)
    assert sheet["K5"].value == datetime(2026, 10, 18)
    assert sheet["K6"].value == "=LITERAL()" and sheet["K6"].data_type == "s"
    assert sheet["E12"].value == "000123" and sheet["E12"].data_type == "s"
    assert sheet[total].value == sheet[summary].value == 324
    assert sheet[f"H{stock}"].value == 1
    assert sheet[f"I{stock}"].value == sheet[f"J{stock}"].value == "000123"
    for col in ("B", "C", "D", "E", "F", "G", "K", "L", "M"):
        assert sheet[f"{col}{stock}"].value is None
    for cell in (closing, remittance, ack_amount, f"H{signature}", f"L{signature}"):
        assert sheet[cell].value is None
    assert sheet[f"A{signature}"].value == "SAMPLE COLLECTOR"
    assert not any(cell.data_type == "f" for row in sheet for cell in row)
    assert not any("DRAFT" in str(cell.value).upper() for row in sheet for cell in row)


def test_exact_form_never_silently_truncates_or_inserts_rows():
    report = sample()
    report["ranges"] = [deepcopy(report["ranges"][0]) for _ in range(6)]
    wb = openpyxl.load_workbook(generate_daily_rcd_excel(report, "0007"))
    assert wb["REMITTANCE 100"].max_row == 72
    assert wb["REMITTANCE 200"].max_row == 67
    assert wb["REMITTANCE 200"]["E17"].value == "000123"
    report["ranges"].append(deepcopy(report["ranges"][0]))
    with pytest.raises(ValueError, match="No receipts were omitted"):
        generate_daily_rcd_excel(report, "0007")


def test_export_contains_no_macro_hidden_source_notes_or_cached_office_entries():
    output = generate_daily_rcd_excel(sample(), "0007")
    with zipfile.ZipFile(output) as package:
        assert not any(
            "vba" in name.lower() or "externallink" in name.lower()
            for name in package.namelist()
        )
    output.seek(0)
    wb = openpyxl.load_workbook(BytesIO(output.getvalue()))
    assert wb.sheetnames == ["REMITTANCE 100", "REMITTANCE 200", "Receipt Details"]
    assert "collections draft" in wb.properties.title.lower()
    assert any(
        "COLLECTIONS DRAFT" in str(row[0].value) for row in wb["Receipt Details"]
    )
    before = deepcopy(office_layout())
    generate_daily_rcd_excel(sample(), "0008")
    assert office_layout() == before


def test_form_parser_rejects_xml_entities_without_resolving_external_data(monkeypatch):
    from defusedxml.common import DefusedXmlException
    from backend.services import rcd_office_form

    tampered = deepcopy(office_layout())
    tampered["sheets"]["100"]["style_components"]["font"][0] = (
        '<!DOCTYPE font [<!ENTITY injected "untrusted">]>'
        '<font><name val="&injected;"/></font>'
    )
    monkeypatch.setattr(rcd_office_form, "office_layout", lambda: tampered)
    with pytest.raises(DefusedXmlException):
        generate_daily_rcd_excel(sample(), "0007")
