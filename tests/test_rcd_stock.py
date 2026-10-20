"""Invented physical forms only. No live stock, receipts or financial writes."""

from copy import copy, deepcopy
from io import BytesIO
import os
from unittest.mock import patch

import openpyxl
import pytest

from backend.services.daily_rcd_excel import generate_daily_rcd_excel
from utils.rcd_remittance import build_remittance_review
from utils.rcd_stock import build_stock_review


def source():
    return {
        "issues": [],
        "collector": {"id": 1, "name": "SAMPLE COLLECTOR"},
        "report_date": "2026-10-18",
        "source_signature": "a" * 64,
        "stock_booklets": [
            {
                "id": 1,
                "collector_id": 1,
                "form_type": "AF 56",
                "first_or": "000123",
                "last_or": "000172",
                "issued_on": "2026-10-01",
            }
        ],
        "totals": {"cash": "648.00", "basic_cash": "324.00", "sef_cash": "324.00"},
        "receipts": [
            {
                "booklet_id": 1,
                "or_number": "000123",
                "booklet_first": "000123",
                "booklet_last": "000172",
                "payment_ids": [1],
                "cash": "648.00",
                "basic_cash": "324.00",
                "sef_cash": "324.00",
            }
        ],
        "ranges": [
            {
                "form_type": "AF 56",
                "from_or": "000123",
                "to_or": "000123",
                "count": 1,
                "basic_cash": "324.00",
                "sef_cash": "324.00",
            }
        ],
    }


def stock(**changes):
    return {
        "confirmed": True,
        "register_reference": "SAMPLE STOCK LOG PAGE 1",
        "opening_ranges": [{"booklet_id": 1, "from_or": "000123", "to_or": "000172"}],
        "received_booklet_ids": [],
        "cancelled": [],
        **changes,
    }


def cancellation(number="000124", **changes):
    return {
        "or_number": number,
        "status": "SPOILED",
        "reason": "SAMPLE printing damage",
        "reference": "SAMPLE LOG PAGE 2",
        "copies_retained": True,
        **changes,
    }


def test_stock_roll_forward_and_cancelled_zero_cash_do_not_change_payload():
    payload = source()
    before = deepcopy(payload)
    reviewed = build_stock_review(payload, stock(cancelled=[cancellation()]))
    assert reviewed["totals"] == {
        "opening": 50,
        "received": 0,
        "issued": 2,
        "ending": 48,
        "paid": 1,
        "cancelled": 1,
    }
    assert reviewed["rows"][0]["issued"]["from_or"] == "000123"
    assert reviewed["rows"][0]["issued"]["to_or"] == "000124"
    assert reviewed["rows"][0]["ending"]["from_or"] == "000125"
    assert payload == before


def test_opening_and_unused_gaps_are_preserved_not_bridged():
    details = stock(
        opening_ranges=[
            {"booklet_id": 1, "from_or": "123", "to_or": "125"},
            {"booklet_id": 1, "from_or": "127", "to_or": "172"},
        ]
    )
    result = build_stock_review(source(), details)
    assert len(result["rows"]) == 2
    assert result["rows"][0]["ending"]["to_or"] == "000125"
    assert result["rows"][1]["ending"]["from_or"] == "000127"
    assert result["totals"]["ending"] == 48


def test_same_day_received_booklet_is_fifty_not_an_opening_balance():
    payload = source()
    payload["stock_booklets"][0]["issued_on"] = "2026-10-18"
    result = build_stock_review(
        payload, stock(opening_ranges=[], received_booklet_ids=[1])
    )
    assert result["totals"]["received"] == 50
    assert result["totals"]["opening"] == 0
    with pytest.raises(ValueError, match="Opening stock"):
        build_stock_review(payload, stock())
    with pytest.raises(ValueError, match="Received stock"):
        build_stock_review(
            payload, stock(opening_ranges=[], received_booklet_ids=[1, 1])
        )


def test_every_printed_row_balances_for_all_ten_serial_usage_patterns():
    # Exhaust all 1,024 combinations, including unused-before-used, paid gaps,
    # no issues and exhausted stock. No category may gain or lose a serial.
    for mask in range(1 << 10):
        payload = source()
        used = {123 + n for n in range(10) if mask & (1 << n)}
        payload["receipts"] = [
            {**payload["receipts"][0], "or_number": str(n)} for n in sorted(used)
        ]
        details = stock(
            opening_ranges=[{"booklet_id": 1, "from_or": "123", "to_or": "132"}]
        )
        result = build_stock_review(payload, details)
        seen = {key: set() for key in ("opening", "received", "issued", "ending")}
        for row in result["rows"]:
            counts = {}
            for key in seen:
                segment = row[key]
                serials = (
                    set(range(int(segment["from_or"]), int(segment["to_or"]) + 1))
                    if segment
                    else set()
                )
                assert not (seen[key] & serials)
                seen[key].update(serials)
                counts[key] = len(serials)
                if segment:
                    assert segment["count"] == len(serials)
            assert (
                counts["opening"] + counts["received"]
                == counts["issued"] + counts["ending"]
            )
        assert seen["opening"] == set(range(123, 133))
        assert seen["issued"] == used
        assert seen["ending"] == set(range(123, 133)) - used


def test_idle_unused_booklets_are_included_only_by_office_confirmation():
    payload = source()
    payload["stock_booklets"].append(
        {
            **payload["stock_booklets"][0],
            "id": 2,
            "first_or": "000223",
            "last_or": "000272",
        }
    )
    details = stock()
    details["opening_ranges"].append(
        {"booklet_id": 2, "from_or": "000223", "to_or": "000272"}
    )
    result = build_stock_review(payload, details)
    assert result["totals"]["ending"] == 99
    assert result["rows"][1]["issued"] is None
    assert result["rows"][1]["ending"]["count"] == 50


@pytest.mark.parametrize(
    "changes",
    [
        {"confirmed": False},
        {"confirmed": "true"},
        {"register_reference": ""},
        {"opening_ranges": []},
        {"received_booklet_ids": [1]},
        {"opening_ranges": [{"booklet_id": True, "from_or": "123", "to_or": "172"}]},
        {"opening_ranges": [{"booklet_id": 2, "from_or": "123", "to_or": "172"}]},
        {"opening_ranges": [{"booklet_id": 1, "from_or": "122", "to_or": "172"}]},
        {"opening_ranges": [{"booklet_id": 1, "from_or": "A-123", "to_or": "A-172"}]},
        {"opening_ranges": [{"booklet_id": 1, "from_or": "123", "to_or": "172"}] * 2},
        {"cancelled": [cancellation("000123")]},
        {"cancelled": [cancellation("000173")]},
        {"cancelled": [cancellation(), cancellation("124")]},
        {"cancelled": [cancellation(copies_retained=False)]},
        {"cancelled": [cancellation(reason="")]},
        {"cancelled": [cancellation(reference="")]},
        {"cancelled": [cancellation(status="DELETED")]},
        {"cancelled": [cancellation(reason="a\nb")]},
    ],
)
def test_invalid_stock_or_payment_cancellation_is_blocked(changes):
    with pytest.raises(ValueError):
        build_stock_review(source(), stock(**changes))


def test_missing_stock_is_not_silently_filled_and_unknown_fields_are_rejected():
    with pytest.raises(ValueError):
        build_stock_review(source(), {})
    with pytest.raises(ValueError):
        build_stock_review(source(), {**stock(), "destroyed": True})
    payload = source()
    payload["stock_booklets"][0]["collector_id"] = 99
    with pytest.raises(ValueError):
        build_stock_review(payload, stock())


def test_too_many_stock_fragments_block_instead_of_stretching_ranges():
    # One continuous opening, but seven separated spoiled ORs fragment ending.
    details = stock(cancelled=[cancellation(str(n)) for n in range(125, 138, 2)])
    with pytest.raises(ValueError, match="more than six"):
        build_stock_review(source(), details)


def test_stock_declarations_change_review_hash_and_normalize_padding():
    a = build_remittance_review(
        source(), {"cash_check_confirmed": True, "stock": stock()}
    )
    b = build_remittance_review(
        source(),
        {"cash_check_confirmed": True, "stock": stock(cancelled=[cancellation("124")])},
    )
    assert a["review_signature"] != b["review_signature"]
    assert b["details"]["stock"]["cancelled"][0]["or_number"] == "000124"
    assert a["funds"] == b["funds"]


def test_exact_form_stock_slots_zero_cancellation_and_blank_receiver_ack():
    details = {
        "cash_check_confirmed": True,
        "stock": stock(cancelled=[cancellation(reason="=SAMPLE()")]),
    }
    wb = openpyxl.load_workbook(
        BytesIO(generate_daily_rcd_excel(source(), "0007", review=details).getvalue())
    )
    baseline = openpyxl.load_workbook(generate_daily_rcd_excel(source(), "0007"))
    for fund, row, total, ack in (("100", 37, "K21", "K68"), ("200", 34, "K18", "K63")):
        sheet = wb[f"REMITTANCE {fund}"]
        original = baseline[f"REMITTANCE {fund}"]
        assert sheet.print_area == original.print_area
        assert sheet.page_setup == original.page_setup
        assert sheet.page_margins == original.page_margins
        assert set(sheet.merged_cells.ranges) == set(original.merged_cells.ranges)
        for cell_row in sheet:
            for cell in cell_row:
                other = original.cell(cell.row, cell.column)
                for kind in ("font", "fill", "border", "alignment", "protection"):
                    assert copy(getattr(cell, kind)) == copy(getattr(other, kind))
        assert [sheet[f"{col}{row}"].value for col in "BCDEHIJKLM"] == [
            50,
            "000123",
            "000172",
            0,
            2,
            "000123",
            "000124",
            48,
            "000125",
            "000172",
        ]
        assert sheet[total].value == 324
        assert sheet[ack].value is None
        assert sheet[f"C{row}"].data_type == "s"
    supporting = wb["Receipt Details"]
    records = [r for r in supporting.values if r[0] == "000124"]
    assert len(records) == 1 and records[0][4] == 0
    reason_cells = [c for row in supporting for c in row if c.value == "=SAMPLE()"]
    assert reason_cells[0].data_type == "s"


@pytest.mark.skipif(os.name != "nt", reason="Invisible Windows Tk stock review")
def test_stock_ui_does_not_guess_opening_and_resets_confirmation_after_edits():
    import customtkinter as ctk
    from ui.rcd_stock import RcdStockReview

    root = ctk.CTk()
    root.withdraw()
    root.attributes("-alpha", 0)
    view = RcdStockReview(root, source())
    try:
        assert view.values()["confirmed"] is False
        assert view.values()["opening_ranges"] == []
        view.first.insert(0, "000123")
        view.last.insert(0, "000172")
        view.confirmed.set(True)
        with patch("ui.rcd_stock.messagebox.showerror") as opening_error:
            view._add_opening()
        assert not opening_error.called, opening_error.call_args
        assert view.confirmed.get() is False
        assert view.values()["opening_ranges"][0]["from_or"] == "000123"
        view.cancel_or.insert(0, "000124")
        view.reason.insert(0, "SAMPLE damage")
        view.cancel_reference.insert(0, "SAMPLE LOG 1")
        with patch("ui.rcd_stock.messagebox.showerror") as error:
            view._add_cancelled()
        assert error.called and not view.cancelled
        view.copies.set(True)
        view._add_cancelled()
        assert len(view.cancelled) == 1 and view.copies.get() is False
        view.stock_tree.selection_set("o0")
        view._remove_stock()
        assert not view.opening
    finally:
        view.destroy()
        root.destroy()
