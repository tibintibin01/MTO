"""Regression cases use synthetic records only, never the office database."""

import io
import asyncio
import sys
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, patch
from datetime import datetime
from decimal import Decimal

import openpyxl
import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from backend.database import Base
from backend.models import Payment, PaymentBilling, Property, PropertyBilling, TaxPolicy
from backend.services import billing_service as billing
from backend.services import property_service as properties
from backend.services.coa_rcd_generator import generate_coa_rcd_excel
from backend.services.payment_adjustments import (
    equal_adjustment_shares,
    allocated_adjustment_expr,
)
from backend.services.report_payload_service import build_receivables_report
from backend.services.report_accuracy_impact_service import (
    review_report_accuracy_impact,
)
from backend.routes import billing as routes

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "clients" / "desktop"))
from ui.reports import ReportsPage
import ui.reports as reports_ui


@pytest.fixture
def db():
    engine = create_engine(
        "sqlite://", poolclass=StaticPool, connect_args={"check_same_thread": False}
    )
    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine)()
    try:
        yield session
    finally:
        session.rollback()
        session.close()
        engine.dispose()


def account(db, key, year=2026, barangay="SYNTHETIC", assessed=100000):
    prop = Property(
        td_number=f"DEMO-{key}",
        owner_name="SYNTHETIC OWNER",
        barangay=barangay,
        assessed_value=assessed,
        effectivity_date="2023",
    )
    db.add(prop)
    db.flush()
    row = PropertyBilling(
        property_id=prop.id,
        tax_year=year,
        assessed_value=assessed,
        amount_paid=0,
        penalty=0,
        discount=0,
    )
    db.add(row)
    db.flush()
    return prop, row


def receipt(db, prop, row, amount, paid_on, penalty=0, discount=0, linked=True):
    payment = Payment(
        property_id=prop.id,
        amount=amount,
        penalty=penalty,
        discount=discount,
        date_paid=paid_on,
        tax_year=str(row.tax_year),
        or_number=f"DEMO-{prop.id}-{row.id}",
        posted_by="DEMO OPERATOR",
    )
    db.add(payment)
    db.flush()
    if linked:
        db.add(
            PaymentBilling(
                payment_id=payment.id,
                billing_id=row.id,
                tax_year=row.tax_year,
                amount_paid=amount,
            )
        )
        row.amount_paid, row.penalty, row.discount = amount, penalty, discount
    db.commit()
    return payment


def rcd(db):
    return openpyxl.load_workbook(
        generate_coa_rcd_excel(db, datetime(2026, 1, 1), datetime(2026, 12, 31)),
        data_only=True,
    )


@pytest.mark.parametrize("cash,penalty,discount", [(2476.32, 72.12, 0), (1800, 0, 200)])
def test_rcd_cash_matches_actual_posting_path(db, cash, penalty, discount):
    assessed = (
        Decimal(str(cash)) - Decimal(str(penalty)) + Decimal(str(discount))
    ) / Decimal("0.02")
    prop, _ = account(db, "CASH", assessed=assessed)
    properties._sync_financial_records(
        prop.id,
        {
            "TD Number": prop.td_number,
            "Owner Name": prop.owner_name,
            "Assessed Value": str(assessed),
            "Tax Year": "2026",
            "OR Number": "DEMO-CASH",
            "OR Date": "2026-10-05",
            "Penalty": str(penalty),
            "Discount": str(discount),
            "Amount Paid": str(cash),
        },
        db,
    )
    db.commit()
    workbook = rcd(db)
    assert workbook["Summary"]["C18"].value == pytest.approx(cash)
    assert workbook["Detailed Collections"]["H5"].value == pytest.approx(cash)
    assert sum(
        workbook["Summary"].cell(row, 3).value for row in range(10, 18)
    ) == pytest.approx(cash)


@pytest.mark.parametrize("penalty,discount,cash", [(200, 0, 4200), (0, 200, 3800)])
def test_multi_year_adjustments_are_not_repeated(db, penalty, discount, cash):
    prop, _ = account(db, "MULTI", year=2025)
    properties._sync_financial_records(
        prop.id,
        {
            "TD Number": prop.td_number,
            "Owner Name": prop.owner_name,
            "Assessed Value": "100000",
            "Tax Year": "2025,2026",
            "OR Number": "DEMO-MULTI",
            "OR Date": "2026-10-05",
            "Penalty": str(penalty),
            "Discount": str(discount),
            "Amount Paid": str(cash),
        },
        db,
    )
    db.commit()
    rows = db.query(PropertyBilling).order_by(PropertyBilling.tax_year).all()
    assert sum(float(row.penalty) for row in rows) == penalty
    assert sum(float(row.discount) for row in rows) == discount
    summary = billing.get_rpt_receivables_summary(2026, db_session=db)
    assert summary["ending_receivable"] == pytest.approx(0)
    assert summary["equation_variance"] == pytest.approx(0)
    workbook = rcd(db)
    assert workbook["Summary"]["C18"].value == cash
    assert workbook["Summary"]["C10"].value == 1000
    assert workbook["Summary"]["C11"].value == 1000


@pytest.mark.parametrize(
    "adjustment,count", [("0.01", 3), ("0.02", 3), ("200.01", 7), ("0.05", 2)]
)
def test_sql_and_python_adjustment_splits_conserve_every_cent(db, adjustment, count):
    prop, first = account(db, "ROUND", year=2020)
    payment = receipt(
        db, prop, first, count, datetime(2026, 1, 1), penalty=adjustment, linked=False
    )
    rows = [first]
    for year in range(2021, 2020 + count):
        row = PropertyBilling(
            property_id=prop.id,
            tax_year=year,
            assessed_value=100000,
            amount_paid=0,
            penalty=0,
            discount=0,
        )
        db.add(row)
        db.flush()
        rows.append(row)
    # Insert backwards: remainder ordering is tax-year/id, not insertion order.
    for row in reversed(rows):
        db.add(
            PaymentBilling(
                payment_id=payment.id,
                billing_id=row.id,
                tax_year=row.tax_year,
                amount_paid=1,
            )
        )
    db.commit()
    shares = (
        db.query(allocated_adjustment_expr(db, Payment.penalty))
        .select_from(PaymentBilling)
        .join(Payment, Payment.id == PaymentBilling.payment_id)
        .order_by(PaymentBilling.tax_year)
        .all()
    )
    actual = [Decimal(str(value[0])).quantize(Decimal("0.01")) for value in shares]
    assert actual == equal_adjustment_shares(adjustment, count)
    assert sum(actual) == Decimal(adjustment)


def test_historical_balances_exclude_later_penalties(db):
    prop, row = account(db, "HISTORY", year=2023)
    db.commit()
    before = billing.get_rpt_receivables_summary(2024, db_session=db)
    receipt(db, prop, row, 2140, datetime(2025, 1, 20), penalty=140)
    after = billing.get_rpt_receivables_summary(2024, db_session=db)
    assert after["beginning_receivable"] == before["beginning_receivable"] == 2000
    assert after["ending_receivable"] == before["ending_receivable"] == 2000
    posted = billing.get_rpt_receivables_summary(2025, db_session=db)
    assert posted["beginning_receivable"] == 2000
    assert posted["prior_year_adjustments"] == 140
    assert posted["ending_receivable"] == 0
    assert posted["equation_variance"] == 0


def test_credit_is_separate_from_other_accounts_receivable(db):
    prop, row = account(db, "CREDIT")
    account(db, "UNPAID")
    receipt(db, prop, row, 2500, datetime(2026, 5, 1))
    summary = billing.get_rpt_receivables_summary(2026, db_session=db)
    assert summary["ending_receivable"] == 2000
    assert summary["ending_credit"] == 500
    assert summary["credit_movement"] == 500
    assert summary["equation_variance"] == 0


def test_reconciliation_cash_includes_unlinked_receipts(db):
    prop, row = account(db, "UNLINKED")
    receipt(db, prop, row, 500, datetime(2026, 5, 1), linked=False)
    metrics = billing.get_reconciliation_metrics(2026, db_session=db)
    assert metrics["treasury"]["cash_collected_this_year"] == 500
    assert metrics["treasury"]["calendar_applicable_collections"] == 0


def test_rcd_uses_year_specific_configured_rates(db):
    prop, row = account(db, "RATES")
    db.add(
        TaxPolicy(tax_year=2026, basic_rate=Decimal("0.015"), sef_rate=Decimal("0.005"))
    )
    receipt(db, prop, row, 2000, datetime(2026, 5, 1))
    workbook = rcd(db)
    assert workbook["Summary"]["C10"].value == 1500
    assert workbook["Summary"]["C14"].value == 500
    assert workbook["Summary"]["C18"].value == 2000


def test_rcd_blocks_ambiguous_multi_year_receipt_without_links(db):
    prop, row = account(db, "AMBIGUOUS")
    payment = receipt(db, prop, row, 4000, datetime(2026, 5, 1), linked=False)
    payment.tax_year = "2025,2026"
    db.commit()
    with pytest.raises(ValueError, match="no billing allocations"):
        rcd(db)


def test_fiscal_barangay_cutoff_and_credit_totals_match_summary(db):
    prop, row = account(db, "FUTURE", year=2025)
    receipt(db, prop, row, 2000, datetime(2026, 1, 20))
    report = build_receivables_report(2025, db)
    assert report["barangays"][0][6] == report["summary"]["ending_receivable"] == 2000
    # Keep the older, explicitly billing-year service default compatible.
    assert properties.get_receivables_by_barangay(2025, db_session=db)[0][6] == 0

    credit, credit_billing = account(db, "BRGY-CREDIT", year=2025)
    receipt(db, credit, credit_billing, 2500, datetime(2025, 5, 1))
    report = build_receivables_report(2025, db)
    assert report["barangays"][0][6] == report["summary"]["ending_receivable"] == 2000
    assert report["barangays"][0][7] == report["summary"]["ending_credit"] == 500


def test_barangay_ranking_does_not_deduct_discount_twice(db):
    a, ba = account(db, "RANK-A", barangay="A")
    account(db, "RANK-B", barangay="B", assessed=90000)
    receipt(db, a, ba, 100, datetime(2026, 5, 1), discount=100)
    # A = 2000 - 100 discount - 100 cash = 1800; B = 1800.
    rows = properties.get_receivables_by_barangay(2026, db_session=db)
    assert [row[0] for row in rows] == ["A", "B"]
    assert [row[6] for row in rows] == [1800, 1800]


def test_receivables_payload_does_not_update_financial_records(db):
    prop, row = account(db, "NO-WRITES")
    receipt(db, prop, row, 500, datetime(2026, 5, 1))
    before = (row.amount_paid, row.penalty, row.discount, db.query(Payment).count())
    report = build_receivables_report(2026, db, diagnostics=True, full_details=True)
    db.refresh(row)
    assert before == (
        row.amount_paid,
        row.penalty,
        row.discount,
        db.query(Payment).count(),
    )
    assert report["diagnostics"]["details_complete"] is True
    assert (
        sum(row[6] for row in report["barangays"])
        == report["summary"]["ending_receivable"]
    )


async def exported_workbook(db, report_type, year="2026"):
    reply = await routes.export_billing_excel(
        routes.ExportReportRequest(report_type=report_type, year=year),
        current_user={"id": 0, "username": "SYNTHETIC", "role": "admin"},
        db_session=db,
    )
    chunks = [chunk async for chunk in reply.body_iterator]
    return openpyxl.load_workbook(io.BytesIO(b"".join(chunks)), data_only=False)


def test_actual_collection_export_retains_all_502_rows_and_literal_usernames(db):
    prop, row = account(db, "EXPORT")
    for index in range(502):
        db.add(
            Payment(
                property_id=prop.id,
                amount=Decimal("1.01"),
                tax_year="2026",
                date_paid=datetime(2026, 5, 1),
                or_number=f"DEMO-EXPORT-{index}",
                posted_by="=DEMO OPERATOR",
            )
        )
    db.commit()
    workbook = asyncio.run(exported_workbook(db, "collections"))
    sheet = workbook.active
    assert sheet.max_row == 507  # title/header + 502 records + total
    assert sheet["G507"].value == pytest.approx(507.02)
    assert sheet["H5"].value == "=DEMO OPERATOR"
    assert sheet["H5"].data_type == "s"
    assert isinstance(sheet["A5"].value, datetime)
    assert [sheet.cell(4, col).value for col in range(1, 9)] == [
        "Date Paid",
        "OR Number",
        "TD Number",
        "Owner Name",
        "Property Kind",
        "Tax Year",
        "Cash Received",
        "Posted By",
    ]


def test_actual_summary_export_includes_prepaid_credit_and_balances(db):
    prop, row = account(db, "PREPAID-EXPORT")
    receipt(db, prop, row, 2000, datetime(2025, 12, 20))
    sheet = asyncio.run(exported_workbook(db, "receivables")).active
    values = {
        sheet.cell(r, 1).value: sheet.cell(r, 2).value
        for r in range(5, sheet.max_row + 1)
    }
    assert values["Less: Prepaid Credits Applied"] == 2000
    assert values["Ending Receivable"] == values["Equation Variance"] == 0
    computed = (
        values["Beginning Receivable"]
        + values["Current-Year Charge (levy + penalty - discount)"]
        - values["Less: Calendar-Year Collections Applied"]
        - values["Less: Prepaid Credits Applied"]
        + values["Prior-Year Penalty/Discount Movements"]
        + values["Credit Movement (separate from debt)"]
    )
    assert computed == values["Ending Receivable"]


def test_actual_barangay_export_keeps_credits_separate(db):
    prop, row = account(db, "BRGY-EXPORT")
    account(db, "BRGY-UNPAID")
    receipt(db, prop, row, 2500, datetime(2026, 5, 1))
    sheet = asyncio.run(exported_workbook(db, "receivables_by_barangay")).active
    assert sheet["G4"].value == "Positive Receivable"
    assert sheet["H4"].value == "Separate Credit"
    assert sheet["G5"].value == sheet["G6"].value == 2000
    assert sheet["H5"].value == sheet["H6"].value == 500


def test_unlinked_cash_never_has_healthy_reconciliation_headline(db):
    prop, row = account(db, "STATUS")
    receipt(db, prop, row, 500, datetime(2026, 5, 1), linked=False)
    report = build_receivables_report(2026, db, diagnostics=True, full_details=True)
    page = ReportsPage.__new__(ReportsPage)
    page.recon_year_cb = SimpleNamespace(get=lambda: "2026")
    payload = page._build_reconciliation_payload(
        report["summary"], report["barangays"], report["metrics"], report["diagnostics"]
    )
    assert payload["status"] == "NEEDS REVIEW"
    assert payload["treasury"]["cash_collected_this_year"] == 500
    assert payload["treasury"]["cash_allocation_difference"] == 500


def test_offsetting_receipt_allocation_errors_do_not_cancel_diagnostics(db):
    for key, allocated in [("SHORT", 400), ("LONG", 600)]:
        prop, row = account(db, key)
        payment = receipt(db, prop, row, 500, datetime(2026, 5, 1), linked=False)
        db.add(
            PaymentBilling(
                payment_id=payment.id,
                billing_id=row.id,
                tax_year=2026,
                amount_paid=allocated,
            )
        )
        row.amount_paid = allocated
    db.commit()
    diagnostics = billing.get_reconciliation_diagnostics(2026, db_session=db)
    assert diagnostics["payment_link_issue_count"] == 2
    assert len(diagnostics["receipt_allocation_mismatches"]) == 2


def test_full_diagnostics_include_more_than_200_rows(db):
    prop, row = account(db, "MANY-ISSUES")
    for index in range(205):
        db.add(
            Payment(
                property_id=prop.id,
                amount=1,
                tax_year="2026",
                date_paid=datetime(2026, 5, 1),
                or_number=f"DEMO-ISSUE-{index}",
            )
        )
    db.commit()
    limited = billing.get_reconciliation_diagnostics(2026, limit=25, db_session=db)
    full = billing.get_reconciliation_diagnostics(2026, limit=None, db_session=db)
    assert limited["unlinked_payment_count"] == full["unlinked_payment_count"] == 205
    assert len(limited["unlinked_payments"]) == 25
    assert len(full["unlinked_payments"]) == 205


def test_collection_stale_request_is_discarded_and_pagination_is_not_reentrant():
    page = ReportsPage.__new__(ReportsPage)
    selected = {"month": "01"}
    page.month_cb = SimpleNamespace(get=lambda: selected["month"])
    page.year_cb = SimpleNamespace(get=lambda: "2026")
    page._coll_cursors, page._coll_page, page._coll_page_size = [None], 0, 100
    page._coll_has_more = True
    page._show_loading = page._hide_loading = lambda: None
    callbacks, workers, displayed = [], [], []
    page.container = SimpleNamespace(
        after=lambda delay, callback: callbacks.append(callback)
    )
    page._update_coll_table = lambda data: displayed.append(data["month"])

    class DeferredThread:
        def __init__(self, target, **kwargs):
            self.target = target

        def start(self):
            workers.append(self.target)

    with (
        patch.object(reports_ui.threading, "Thread", DeferredThread),
        patch.object(
            reports_ui.billing,
            "get_report_details",
            side_effect=lambda month, *args, **kwargs: {"month": month},
        ),
    ):
        page._coll_load_page()
        page._coll_next_page()
        assert (
            page._coll_page == 0
        )  # Clicking next during a pending load cannot skip a cursor.
        selected["month"] = "02"
        page._coll_load_page()
        workers[1]()
        for callback in callbacks:
            callback()
        callbacks.clear()
        workers[0]()
        for callback in callbacks:
            callback()
    assert displayed == ["02"]
    assert not page._report_pending


def test_all_year_rcd_range_is_not_replaced_with_this_year():
    page = ReportsPage.__new__(ReportsPage)
    page.month_cb = SimpleNamespace(get=lambda: "All")
    page.year_cb = SimpleNamespace(get=lambda: "All")
    assert page.get_selected_date_range() == ("1000-01-01", "9999-12-31")
    page.month_cb = SimpleNamespace(get=lambda: "02")
    with pytest.raises(ValueError, match="Select a year"):
        page.get_selected_date_range()


def test_aggregate_impact_review_is_private_and_does_not_repair_cache(db):
    prop, row = account(db, "IMPACT")
    receipt(db, prop, row, 2140, datetime(2026, 5, 1), penalty=140)
    row.penalty = 280  # Synthetic legacy drift; the review must leave it untouched.
    db.commit()
    result = review_report_accuracy_impact(db)
    db.refresh(row)
    assert row.penalty == 280
    assert result["counts"]["cached_billing_adjustment_mismatches"] == 1
    assert result["financial_records_changed"] is False
    assert "SYNTHETIC OWNER" not in str(result)
    assert prop.td_number not in str(result)


def test_summary_screen_explains_prepaid_credits(db):
    prop, row = account(db, "PREPAID-SCREEN")
    receipt(db, prop, row, 2000, datetime(2025, 12, 20))
    summary = billing.get_rpt_receivables_summary(2026, db_session=db)
    page = ReportsPage.__new__(ReportsPage)
    page.receiv_content = MagicMock()
    page.receiv_content.winfo_children.return_value = []
    labels = []

    def label(*args, **kwargs):
        labels.append(kwargs.get("text", ""))
        return MagicMock()

    with (
        patch.object(reports_ui.ctk, "CTkFrame", return_value=MagicMock()),
        patch.object(reports_ui.ctk, "CTkLabel", side_effect=label),
        patch.object(reports_ui.ctk, "CTkProgressBar", return_value=MagicMock()),
        patch.object(reports_ui.ttk, "Treeview", return_value=MagicMock()),
    ):
        page._update_receiv_summary(summary, [])
    equation = next(text for text in labels if "ending receivable" in text)
    assert "2,000.00 prepaid credits" in equation
    assert any("Coverage: 100.0%" in text for text in labels)
    assert not any("Remaining target: 2,000" in text for text in labels)


def test_reconciliation_pdf_sends_all_diagnostics_to_actual_renderer(db, tmp_path):
    prop, row = account(db, "PDF")
    receipt(db, prop, row, 2000, datetime(2026, 5, 1))
    report = build_receivables_report(2026, db, diagnostics=True, full_details=True)
    page = ReportsPage.__new__(ReportsPage)
    page.recon_year_cb = SimpleNamespace(get=lambda: "2026")
    payload = page._build_reconciliation_payload(
        report["summary"], report["barangays"], report["metrics"], report["diagnostics"]
    )
    payload["diagnostic_rows"] = [
        ("Synthetic issue", f"DEMO-DIAG-{i:03}", 2026, "DEMO", 1) for i in range(1, 42)
    ]
    import reportlab.platypus as layout

    real_table = layout.Table
    sent_rows = []

    def table(rows, *args, **kwargs):
        sent_rows.extend(rows)
        return real_table(rows, *args, **kwargs)

    destination = tmp_path / "reconciliation-complete.pdf"
    with patch.object(layout, "Table", side_effect=table):
        page._write_reconciliation_pdf(payload, str(destination))
    assert destination.read_bytes().startswith(b"%PDF")
    text = " ".join(str(cell) for row in sent_rows for cell in row)
    assert "DEMO-DIAG-041" in text


def test_rcd_keeps_taxpayer_text_literal_not_a_formula(db):
    prop, row = account(db, "TEXT")
    prop.owner_name = '=HYPERLINK("https://example.invalid","text")'
    receipt(db, prop, row, 2000, datetime(2026, 5, 1))
    workbook = openpyxl.load_workbook(
        generate_coa_rcd_excel(db, datetime(2026, 1, 1), datetime(2026, 12, 31)),
        data_only=False,
    )
    assert workbook["Detailed Collections"]["C5"].value == prop.owner_name
    assert workbook["Detailed Collections"]["C5"].data_type == "s"
    assert workbook["UACS Summary"]["C13"].value == "=SUM(C5:C12)"
    assert workbook["UACS Summary"]["C13"].data_type == "f"
    detailed = workbook["Detailed Collections"]
    assert detailed["H6"].value == "=SUM(H5:H5)"
    assert detailed["H6"].data_type == "f"


@pytest.mark.parametrize("penalty,discount", [(100, 0), (0, 100)])
def test_adjusted_receipt_does_not_create_charges_for_unallocated_years(
    db, penalty, discount
):
    from fastapi import HTTPException

    prop, row = account(db, "UNALLOCATED", year=2025)
    db.commit()
    with pytest.raises(HTTPException) as error:
        properties._sync_financial_records(
            prop.id,
            {
                "TD Number": prop.td_number,
                "Owner Name": prop.owner_name,
                "Assessed Value": "100000",
                "Tax Year": "2025,2026",
                "OR Number": "DEMO-PARTIAL-ADJUSTED",
                "OR Date": "2026-10-05",
                "Penalty": str(penalty),
                "Discount": str(discount),
                "Amount Paid": "1000",
            },
            db,
        )
    assert error.value.status_code == 422
    db.rollback()
    assert db.query(Payment).count() == 0
    assert db.query(PropertyBilling).count() == 1
    db.refresh(row)
    assert row.penalty == row.discount == 0


def test_undated_stored_adjustment_is_flagged_not_silently_certified(db):
    prop, row = account(db, "UNDATED", year=2025)
    row.penalty = 140
    db.commit()
    report = build_receivables_report(2026, db, diagnostics=True, full_details=True)
    assert report["diagnostics"]["undated_adjustment_count"] == 1
    page = ReportsPage.__new__(ReportsPage)
    page.recon_year_cb = SimpleNamespace(get=lambda: "2026")
    payload = page._build_reconciliation_payload(
        report["summary"], report["barangays"], report["metrics"], report["diagnostics"]
    )
    assert payload["status"] == "NEEDS REVIEW"
    db.refresh(row)
    assert row.penalty == 140


def test_allocated_receipt_without_date_requires_review(db):
    prop, row = account(db, "NO-DATE")
    receipt(db, prop, row, 1000, None)
    report = build_receivables_report(2026, db, diagnostics=True, full_details=True)
    assert report["diagnostics"]["missing_receipt_date_count"] == 1
    page = ReportsPage.__new__(ReportsPage)
    page.recon_year_cb = SimpleNamespace(get=lambda: "2026")
    payload = page._build_reconciliation_payload(
        report["summary"], report["barangays"], report["metrics"], report["diagnostics"]
    )
    assert payload["status"] == "NEEDS REVIEW"


def test_empty_rcd_has_zero_totals_not_circular_formulas(db):
    workbook = openpyxl.load_workbook(
        generate_coa_rcd_excel(db, datetime(2026, 1, 1), datetime(2026, 12, 31)),
        data_only=False,
    )
    assert workbook["Summary"]["C18"].value == 0
    for column in ["D", "E", "F", "G", "H"]:
        assert workbook["Detailed Collections"][column + "5"].value == 0
        assert workbook["Detailed Collections"][column + "5"].data_type == "n"


def test_closing_property_does_not_erase_genuine_collection_cash(db):
    prop, row = account(db, "CLOSED")
    receipt(db, prop, row, 2000, datetime(2026, 5, 1))
    prop.deleted_at = datetime(2026, 6, 1)
    db.commit()
    ledger = billing.get_report_details("All", "2026", db_session=db)
    assert len(ledger["items"]) == 1
    assert float(ledger["items"][0][6]) == 2000
    assert rcd(db)["Summary"]["C18"].value == 2000
    metrics = billing.get_reconciliation_metrics(2026, db_session=db)
    assert metrics["treasury"]["cash_collected_this_year"] == 2000
    assert metrics["treasury"]["closed_account_cash_this_year"] == 2000
    assert metrics["treasury"]["cash_allocation_difference"] == 0
    assert metrics["treasury"]["calendar_applicable_collections"] == 0
