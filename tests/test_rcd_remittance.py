"""Synthetic cash/check and next-morning deposit plans, not office records."""

from copy import deepcopy
from datetime import date, datetime
from decimal import Decimal
from unittest.mock import patch
from types import SimpleNamespace
import os

import openpyxl
import pytest

from api_clients import reports_service
from backend.services.daily_rcd_excel import generate_daily_rcd_excel
from utils.rcd_remittance import build_remittance_review


def payload():
    return {
        "issues": [],
        "collector": {"id": 1, "name": "SAMPLE COLLECTOR"},
        "report_date": "2026-10-18",
        "source_signature": "a" * 64,
        "stock_booklets": [
            {
                "id": 1,
                "collector_id": 1,
                "form_type": "AF 56",
                "first_or": "000123",
                "last_or": "000172",
                "issued_on": "2026-10-01",
            }
        ],
        "totals": {"basic_cash": "324.00", "sef_cash": "324.00", "cash": "648.00"},
        "ranges": [
            {
                "form_type": "AF 56",
                "from_or": "000123",
                "to_or": "000123",
                "count": 1,
                "basic_cash": "324.00",
                "sef_cash": "324.00",
            }
        ],
        "receipts": [
            {
                "or_number": "000123",
                "booklet_id": 1,
                "cash": "648.00",
                "basic_cash": "324.00",
                "sef_cash": "324.00",
                "booklet_first": "000123",
                "booklet_last": "000172",
                "payment_ids": [1],
            }
        ],
    }


def check(**changes):
    return {
        "or_number": "000123",
        "bank": "SAMPLE BANK",
        "check_number": "000009",
        "payee": "SAMPLE MUNICIPAL TREASURER",
        "amount": "100.00",
        "face_amount": changes.get("amount", "100.00"),
        **changes,
    }


def reviewed(**changes):
    return {"cash_check_confirmed": True, "checks": [], **changes}


def handover(**changes):
    return reviewed(
        full_handover_confirmed=True,
        receiving_officer="SAMPLE LIQUIDATING OFFICER",
        handover_date="2026-10-19",
        opening_100="0.00",
        opening_200="0.00",
        **changes,
    )


def test_default_cash_is_not_assumed_without_confirmation():
    with pytest.raises(ValueError, match="Confirm the cash/check"):
        build_remittance_review(payload(), {})
    with pytest.raises(ValueError):
        build_remittance_review(payload(), {"cash_check_confirmed": "true"})
    data = build_remittance_review(payload(), reviewed())
    assert data["funds"]["100"]["cash"] == "324.00"
    assert data["funds"]["100"]["checks"] == "0.00"
    assert data["funds"]["100"]["closing"] is None
    assert data["details"]["planned_deposit_date"] == "2026-10-19"
    assert data["details"]["handover_date"] is None


def test_mixed_instruments_reconcile_to_net_receipt_without_changing_it():
    original = payload()
    before = deepcopy(original)
    result = build_remittance_review(original, reviewed(checks=[check()]))
    for fund in ("100", "200"):
        assert result["funds"][fund]["cash"] == "274.00"
        assert result["funds"][fund]["checks"] == "50.00"
    assert original == before


def test_one_check_can_cover_multiple_ors_without_printing_it_twice():
    source = payload()
    source["receipts"].append({**source["receipts"][0], "or_number": "000124"})
    source["totals"] = {"basic_cash": "648.00", "sef_cash": "648.00", "cash": "1296.00"}
    result = build_remittance_review(
        source,
        reviewed(
            checks=[
                check(amount="648.00", face_amount="1296.00"),
                check(or_number="124", amount="648.00", face_amount="1296.00"),
            ]
        ),
    )
    assert len(result["checks"]) == 1
    assert result["checks"][0]["100"] == "648.00"
    assert result["funds"]["200"]["cash"] == "0.00"


def test_unequal_funds_and_cent_rounding_conserve_each_cent():
    source = payload()
    source["receipts"][0].update(cash="0.03", basic_cash="0.02", sef_cash="0.01")
    source["totals"] = {"cash": "0.03", "basic_cash": "0.02", "sef_cash": "0.01"}
    result = build_remittance_review(
        source,
        reviewed(checks=[check(check_number=str(i), amount="0.01") for i in range(3)]),
    )
    assert sum(Decimal(c["100"]) for c in result["checks"]) == Decimal("0.02")
    assert sum(Decimal(c["200"]) for c in result["checks"]) == Decimal("0.01")
    assert all(
        Decimal(c["100"]) >= 0 and Decimal(c["200"]) >= 0 for c in result["checks"]
    )


@pytest.mark.parametrize(
    "value", ["NaN", "Infinity", "-1", "1.001", "649.00", "0", "9999999999999"]
)
def test_invalid_or_excess_check_amount_is_blocked(value):
    with pytest.raises(ValueError):
        build_remittance_review(payload(), reviewed(checks=[check(amount=value)]))


@pytest.mark.parametrize(
    "entries",
    [
        [check(), check()],
        [check(or_number="999999")],
        [check(bank="")],
        [check(), check(or_number="000123", amount="1", payee="DIFFERENT")],
    ],
)
def test_bad_check_identity_is_blocked(entries):
    with pytest.raises(ValueError):
        build_remittance_review(payload(), reviewed(checks=entries))


def test_whole_check_must_be_fully_accounted_for_before_export():
    partial = reviewed(checks=[check(amount="100.00", face_amount="200.00")])
    with pytest.raises(ValueError, match="whole amount"):
        build_remittance_review(payload(), partial)
    # Only the in-progress desktop basket allows adding an unfinished allocation.
    assert build_remittance_review(payload(), partial, allow_incomplete_checks=True)
    with pytest.raises(ValueError):
        build_remittance_review(
            payload(),
            reviewed(checks=[check(amount="100.00", face_amount="50.00")]),
            allow_incomplete_checks=True,
        )


def test_actual_handover_and_deposit_plan_are_distinct_and_not_future_claims():
    data = build_remittance_review(payload(), handover(), today=date(2026, 10, 19))
    assert data["funds"]["100"]["remitted"] == "324.00"
    assert data["funds"]["100"]["closing"] == "0.00"
    with pytest.raises(ValueError, match="actual receiving officer"):
        build_remittance_review(payload(), reviewed(full_handover_confirmed=True))
    with pytest.raises(ValueError):
        build_remittance_review(payload(), handover(), today=date(2026, 10, 18))
    with pytest.raises(ValueError):
        build_remittance_review(payload(), reviewed(receiving_officer="SAMPLE"))
    assert (
        build_remittance_review(payload(), reviewed(planned_deposit_date="2026-10-20"))[
            "details"
        ]["planned_deposit_date"]
        == "2026-10-20"
    )


def test_unknown_opening_is_not_a_zero_closing_even_after_handover():
    values = handover()
    values.pop("opening_100")
    values["opening_200"] = "10.00"
    data = build_remittance_review(payload(), values, today=date(2026, 10, 19))
    assert data["funds"]["100"]["closing"] is None
    assert data["funds"]["200"]["closing"] == "10.00"
    assert data["funds"]["200"]["total_accountability"] == "334.00"


def test_review_signature_is_bound_to_source_and_idempotent():
    result = build_remittance_review(payload(), reviewed(checks=[check()]))
    assert (
        build_remittance_review(payload(), result["details"])["review_signature"]
        == result["review_signature"]
    )
    source = payload()
    source["source_signature"] = "b" * 64
    assert (
        build_remittance_review(source, result["details"])["review_signature"]
        != result["review_signature"]
    )


def test_completed_cash_check_fields_use_original_cells_without_bank_acknowledgment():
    with patch(
        "utils.rcd_remittance.philippine_report_date", return_value=date(2026, 10, 19)
    ):
        book = openpyxl.load_workbook(
            generate_daily_rcd_excel(
                payload(), "0007", review=handover(checks=[check()])
            )
        )
    for fund, total, remit, closing, check_row, summary in (
        ("100", "K21", 29, "E58", 49, 52),
        ("200", "K18", 26, "E53", 44, 47),
    ):
        sheet = book["REMITTANCE " + fund]
        assert sheet["K5"].value == datetime(2026, 10, 18)
        assert sheet[f"K{remit}"].value == 274
        assert sheet[f"K{remit+1}"].value == 50
        assert sheet[f"K{remit+2}"].value == sheet[total].value == 324
        assert sheet[closing].value == 0
        assert sheet[f"E{summary+2}"].value == -324
        assert sheet[f"H{check_row}"].value == "000009"
        assert sheet[f"M{check_row}"].value == 50
        assert sheet["H70" if fund == "100" else "H65"].value is None
        assert sheet["K68" if fund == "100" else "K63"].value is None
    assert book["REMITTANCE 100"].max_row == 72
    assert book["REMITTANCE 200"].max_row == 67


def test_review_body_not_url_contains_bank_and_check_details():
    with patch.object(reports_service, "api_download_file") as download:
        reports_service.download_daily_rcd(
            7, "2026-10-18", "0009", "a" * 64, review=reviewed()
        )
    assert download.call_args.args == ("POST", "/reports/rcd/daily/excel")
    assert "params" not in download.call_args.kwargs
    assert download.call_args.kwargs["json_body"]["review"] == reviewed()


def test_reviewed_amounts_retain_original_office_currency_formats():
    from backend.services.rcd_office_form import office_layout
    from openpyxl.utils.cell import coordinate_to_tuple

    with patch(
        "utils.rcd_remittance.philippine_report_date", return_value=date(2026, 10, 19)
    ):
        book = openpyxl.load_workbook(
            generate_daily_rcd_excel(
                payload(), "0007", review=handover(checks=[check()])
            )
        )
    for fund, cells in (
        ("100", ("K29", "K30", "K31", "E50", "E51", "E52", "E54", "E58", "M49", "M58")),
        ("200", ("K26", "K27", "K28", "E45", "E46", "E47", "E49", "E53", "M44", "M53")),
    ):
        spec = office_layout()["sheets"][fund]
        for address in cells:
            row, col = coordinate_to_tuple(address)
            fmt = spec["styles"][spec["cells"][row - 1][col - 1]][5]
            assert book["REMITTANCE " + fund][address].number_format == (
                "#,##0.00" if fmt == "General" else fmt
            )


def test_download_transport_sends_review_as_json_without_queueing():
    from api_clients import api_helper

    response = SimpleNamespace(status_code=200, raise_for_status=lambda: None)
    with (
        patch.object(api_helper, "_SESSION_TOKEN", None),
        patch.object(api_helper, "get_tls_verification", return_value=True),
        patch.object(api_helper, "record_connection_success"),
        patch.object(api_helper.requests, "request", return_value=response) as request,
        patch.object(
            api_helper, "save_stream_response_to_temp_file", return_value="sample.xlsx"
        ),
    ):
        assert (
            api_helper.api_download_file(
                "POST", "/reports/rcd/daily/excel", json_body={"review": reviewed()}
            )
            == "sample.xlsx"
        )
    assert request.call_args.kwargs["json"] == {"review": reviewed()}
    assert request.call_args.kwargs["params"] is None
    assert request.call_args.kwargs["stream"] is True


@pytest.mark.skipif(os.name != "nt", reason="Invisible Windows Tk review dialog")
def test_review_dialog_requires_cash_confirmation_and_explicit_zero_opening():
    import customtkinter as ctk
    from ui.rcd_remittance import RcdRemittanceDialog

    root = ctk.CTk()
    root.withdraw()
    root.attributes("-alpha", 0)
    dialog = None
    try:
        dialog = RcdRemittanceDialog(root, payload())
        dialog.withdraw()
        dialog.attributes("-alpha", 0)
        dialog.geometry("780x640")
        root.deiconify()
        root.update()
        dialog.deiconify()
        ready = ctk.BooleanVar(master=root, value=False)
        root.after(300, lambda: ready.set(True))
        root.wait_variable(ready)
        assert dialog.cash_confirmed.get() is False
        assert dialog.opening_100.get() == dialog.opening_200.get() == ""
        assert dialog.deposit_date.get() == "2026-10-19"
        assert dialog.handover_date.get() == ""
        assert dialog.recipient.cget("state") == "disabled"
        for widget in (
            dialog.or_number,
            dialog.bank,
            dialog.check_number,
            dialog.payee,
            dialog.amount,
            dialog.face_amount,
        ):
            assert widget.winfo_ismapped(), (dialog.tabs.get(), dialog.winfo_ismapped())
            assert (
                widget.winfo_x() + widget.winfo_width() <= widget.master.winfo_width()
            )
        with patch("ui.rcd_remittance.messagebox.showerror") as error:
            dialog._save()
        assert error.called and dialog.result is None
        dialog._zero_opening()
        dialog.cash_confirmed.set(True)
        dialog.stock.reference.insert(0, "SAMPLE LOG PAGE 1")
        dialog.stock.first.insert(0, "000123")
        dialog.stock.last.insert(0, "000172")
        with patch("ui.rcd_stock.messagebox.showerror") as opening_error:
            dialog.stock._add_opening()
        assert not opening_error.called, opening_error.call_args
        dialog.stock.confirmed.set(True)
        with (
            patch("ui.rcd_remittance.messagebox.askyesno", return_value=True),
            patch("ui.rcd_remittance.messagebox.showerror") as save_error,
        ):
            dialog._save()
        assert not save_error.called, save_error.call_args
        assert dialog.result["opening_100"] == dialog.result["opening_200"] == "0.00"
        assert dialog.result["full_handover_confirmed"] is False
        dialog = None
    finally:
        if dialog is not None:
            dialog.destroy()
        root.destroy()
